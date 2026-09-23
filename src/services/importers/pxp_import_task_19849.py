"""
pxp_import_task_19849.py

Importador PXP do processo 19849.

Objetivo
--------
Processar arquivo XLSX de entrada contendo linhas PXP do Cisco LCI e refletir
as regras de negócio no banco de dados com foco em:

- criação e atualização de tasks Cisco LCI;
- atualização de activities (stages) das tasks Cisco LCI;
- atualização de Use Case e Use Case Exit Criteria;
- geração de histórico e arquivos de falha.

Padrão adotado
--------------
Este importador segue a mesma linha operacional de `pxp_import_task_6702.py`
e `pxp_import_activity_5890.py`:

- lê o arquivo XLSX diretamente;
- processa em chunks;
- remove do arquivo original toda linha lida;
- escreve linhas com erro em um arquivo fixo de falhas;
- gera log texto de execução;
- grava logs funcionais em tbImportLog quando disponível;
- registra histórico em tbTaskRecord.

Observações importantes
-----------------------
1. Task Cisco LCI:
   - task_tasktype_id sempre 22 para novas tasks;
   - busca de task existente por WS considera task types 21 e 22;
   - prioridade para tipo 22.

2. Activities Cisco LCI:
   - activities representam estágios: Onboard, Implement, Use, Engage, Adopt;
   - activities só existem após Opted In no portal Cisco PXP;
   - activity_ws também é exclusivo conceitualmente, mas neste importador
     a vinculação principal das filhas acontece por task_id + activity_name.

3. Moeda:
   - task e activity usam USD.

4. Task e activity consideradas encerradas:
   - status IN (4, 5, 6, 10)

5. Sanitização numérica:
   - task_booking_amount é validado contra o range suportado
     pelo banco para DECIMAL(30,6)
   - quando valor numérico vier inválido ou fora do range:
       * a linha NÃO é rejeitada
       * o campo é gravado como None
       * é gerado log funcional
"""

from __future__ import annotations

import logging
import re
import traceback
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.infrastructure.database.repositories.task_repository import TaskRepository
from src.infrastructure.database.repositories.task_activity_repository import TaskActivityRepository
from src.infrastructure.database.repositories.task_history_repository import TaskHistoryRepository
from src.infrastructure.database.repositories.company_list_name_repository import CompanyListNameRepository

try:
    from src.infrastructure.database.repositories.import_log_repository import ImportLogRepository
except Exception:
    ImportLogRepository = None

try:
    from src.infrastructure.database.repositories.use_case_repository import UseCaseRepository
except Exception:
    UseCaseRepository = None

try:
    from src.infrastructure.database.repositories.import_control_repository import (
        ImportControlRepository,
        ImportStatus,
    )
except Exception:
    ImportControlRepository = None
    ImportStatus = None


# ======================================================================================
# CONFIGURAÇÃO BASE
# ======================================================================================

BASE_STORAGE_PATH = Path("/home/bridgeadoption/storage")
BASE_INPUT_PATH = BASE_STORAGE_PATH / "input"
BASE_OUTPUT_PATH = BASE_STORAGE_PATH / "output"
BASE_LOGS_PATH = BASE_STORAGE_PATH / "logs"

IMPORT_SOURCE = "PxpImportTask19849"

ERROR_EXTRA_COLUMNS = [
    "import_error_message",
    "import_error_column",
    "import_error_value",
    "import_original_row",
    "import_processed_at",
    "import_source",
    "import_original_file",
]

logger = logging.getLogger(__name__)

repo_task = TaskRepository()
repo_activity = TaskActivityRepository()
repo_history = TaskHistoryRepository()
repo_company = CompanyListNameRepository()
repo_log = ImportLogRepository() if ImportLogRepository else None
repo_use_case = UseCaseRepository() if UseCaseRepository else None
repo_import_control = ImportControlRepository() if ImportControlRepository else None


# ======================================================================================
# CONSTANTES DE DOMÍNIO
# ======================================================================================

DEFAULT_TASK_OWNER_ID = 93
CHUNK_SIZE = 1000
IMPORT_CONTROL_PROGRESS_INTERVAL_CHUNKS = 1

TASK_TYPE_LCI_LEGACY = 21
TASK_TYPE_LCI_CURRENT = 22
VALID_EXISTING_TASK_TYPE_IDS = {TASK_TYPE_LCI_LEGACY, TASK_TYPE_LCI_CURRENT}
PRIORITY_TASK_TYPE_ID = TASK_TYPE_LCI_CURRENT

VENDOR_ID_CISCO = 1

TASK_STATUS_OPEN = 1
TASK_STATUS_IN_PROGRESS = 2
TASK_STATUS_ON_HOLD = 3
TASK_STATUS_CANCELLED = 4
TASK_STATUS_DECLINED = 5
TASK_STATUS_EXPIRED = 6
TASK_STATUS_COMPLETED = 10

ACTIVITY_STATUS_OPEN = 1
ACTIVITY_STATUS_IN_PROGRESS = 2
ACTIVITY_STATUS_CANCELLED = 4
ACTIVITY_STATUS_DECLINED = 5
ACTIVITY_STATUS_EXPIRED = 6
ACTIVITY_STATUS_COMPLETED = 10

CLOSED_TASK_STATUSES = {
    TASK_STATUS_CANCELLED,
    TASK_STATUS_DECLINED,
    TASK_STATUS_EXPIRED,
    TASK_STATUS_COMPLETED,
}
CLOSED_ACTIVITY_STATUSES = {
    ACTIVITY_STATUS_CANCELLED,
    ACTIVITY_STATUS_DECLINED,
    ACTIVITY_STATUS_EXPIRED,
    ACTIVITY_STATUS_COMPLETED,
}

TASK_NOT_ELIGIBLE_JUSTIFICATION = "TASK NOT ELIGIBLE BY VENDOR"

REQUIRED_COLUMNS = [
    "Deal WS-ID",
    "CR Party Name",
    "CR Party ID",
    "Track",
    "Sub-Track",
    "Deal Incentive Expiry Date",
    "Booking Date",
    "Booking Amount - Net to Cisco",
    "Stage",
    "Adopt Rebate Opt-In Status",
    "Adopt Rebate Start Date",
    "Stage Completion Date(onboard)",
    "Stage Completion Date (Implement)",
    "Stage Completion Date(Use)",
    "Stage Completion Date(Engage)",
    "Stage Completion Date(Adopt)",
    "Task Details (Onboard)",
    "Task Details (Implement)",
    "Task Details (Use)",
    "Task Details (Engage)",
    "Task Details (Adopt)",
    "Deal ID",
    "Deal CPI Portfolio",
]

OPTIONAL_COLUMNS = [
    "Incentive Level",
    "Stage Completion Flag(onboard)",
    "Stage Completion Flag (Implement)",
    "Stage Completion Flag(Use)",
    "Stage Completion Flag(Engage)",
    "Stage Completion Flag(Adopt)",
    "EA Flag",
    "Subscription ID",
    "Booking PO Number",
]

EXPECTED_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS

HEADER_ALIASES = {
    "Stage Completion Flag(onboard)": [
        "Stage Completion Flag(onboard)",
        "Stage Completion Flag (onboard)",
    ],
    "Stage Completion Flag (Implement)": [
        "Stage Completion Flag (Implement)",
        "Stage Completion Flag(Implement)",
    ],
    "Stage Completion Flag(Use)": [
        "Stage Completion Flag(Use)",
        "Stage Completion Flag (Use)",
    ],
    "Stage Completion Flag(Engage)": [
        "Stage Completion Flag(Engage)",
        "Stage Completion Flag (Engage)",
    ],
    "Stage Completion Flag(Adopt)": [
        "Stage Completion Flag(Adopt)",
        "Stage Completion Flag (Adopt)",
    ],
    "Stage Completion Date(onboard)": [
        "Stage Completion Date(onboard)",
        "Stage Completion Date (onboard)",
    ],
    "Stage Completion Date (Implement)": [
        "Stage Completion Date (Implement)",
        "Stage Completion Date(Implement)",
    ],
    "Stage Completion Date(Use)": [
        "Stage Completion Date(Use)",
        "Stage Completion Date (Use)",
    ],
    "Stage Completion Date(Engage)": [
        "Stage Completion Date(Engage)",
        "Stage Completion Date (Engage)",
    ],
    "Stage Completion Date(Adopt)": [
        "Stage Completion Date(Adopt)",
        "Stage Completion Date (Adopt)",
    ],
    "Task Details (Onboard)": [
        "Task Details (Onboard)",
        "Task Details(Onboard)",
    ],
    "Task Details (Implement)": [
        "Task Details (Implement)",
        "Task Details(Implement)",
    ],
    "Task Details (Use)": [
        "Task Details (Use)",
        "Task Details(Use)",
    ],
    "Task Details (Engage)": [
        "Task Details (Engage)",
        "Task Details(Engage)",
    ],
    "Task Details (Adopt)": [
        "Task Details (Adopt)",
        "Task Details(Adopt)",
    ],
}

STAGE_CONFIG = {
    "Onboard": {
        "details_header": "Task Details (Onboard)",
        "date_header": "Stage Completion Date(onboard)",
        "flag_header": "Stage Completion Flag(onboard)",
        "seq": 1,
    },
    "Implement": {
        "details_header": "Task Details (Implement)",
        "date_header": "Stage Completion Date (Implement)",
        "flag_header": "Stage Completion Flag (Implement)",
        "seq": 2,
    },
    "Use": {
        "details_header": "Task Details (Use)",
        "date_header": "Stage Completion Date(Use)",
        "flag_header": "Stage Completion Flag(Use)",
        "seq": 3,
    },
    "Engage": {
        "details_header": "Task Details (Engage)",
        "date_header": "Stage Completion Date(Engage)",
        "flag_header": "Stage Completion Flag(Engage)",
        "seq": 4,
    },
    "Adopt": {
        "details_header": "Task Details (Adopt)",
        "date_header": "Stage Completion Date(Adopt)",
        "flag_header": "Stage Completion Flag(Adopt)",
        "seq": 5,
    },
}

DECIMAL_30_6_MAX = Decimal("999999999999999999999999.999999")
DECIMAL_30_6_MIN = Decimal("-999999999999999999999999.999999")
DECIMAL_30_6_QUANTIZER = Decimal("0.000001")


# ======================================================================================
# DATACLASSES
# ======================================================================================

@dataclass
class RowProcessResult:
    success: bool
    error_message: Optional[str] = None
    error_column: Optional[str] = None
    error_value: Optional[Any] = None
    ignored: bool = False
    created: bool = False
    updated: bool = False
    cancelled: bool = False
    activities_updated: int = 0
    activities_cancelled: int = 0


# ======================================================================================
# HELPERS DE DIRETÓRIO / LOG
# ======================================================================================

def _ensure_directories() -> None:
    BASE_INPUT_PATH.mkdir(parents=True, exist_ok=True)
    BASE_OUTPUT_PATH.mkdir(parents=True, exist_ok=True)
    BASE_LOGS_PATH.mkdir(parents=True, exist_ok=True)


def _build_execution_log_path(input_path: Path) -> Path:
    return BASE_LOGS_PATH / f"{input_path.stem}.log"


def _report_progress(
    import_control_id: Optional[int],
    total_rows: int,
    processed_success: int,
    failed_rows: int,
    ignored_rows: int,
    remaining_rows: int,
    execution_log_path: Optional[Path] = None,
) -> None:
    if repo_import_control is None or import_control_id is None:
        return

    lidas = total_rows - remaining_rows
    msg = (
        f"Arquivo processado. "
        f"Total linhas origem={total_rows} | "
        f"lidas={lidas} | "
        f"sucesso={processed_success} | "
        f"erros={failed_rows} | "
        f"ignoradas={ignored_rows} | "
        f"restantes={remaining_rows}"
    )

    try:
        repo_import_control.update_status(
            importctrl_id=import_control_id,
            status=ImportStatus.RUNNING,
            message=msg,
        )
    except Exception as e:
        if execution_log_path:
            _append_execution_log(
                execution_log_path,
                f"WARN failed updating import control progress error={str(e)[:500]}",
            )


def _append_execution_log(log_path: Path, message: str) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{ts}] {message}\n")


def _safe_log(
    file_path: str,
    row_number: int,
    message: str,
    column_name: Optional[str] = None,
    value: Optional[Any] = None,
) -> None:
    if repo_log is None:
        return

    try:
        repo_log.create(
            IMPORT_SOURCE,
            file_path,
            row_number,
            message,
            column_name,
            value,
        )
    except Exception as log_ex:
        logger.exception(
            "Falha ao gravar tbImportLog. file=%s row=%s column=%s value=%s error=%s",
            file_path,
            row_number,
            column_name,
            value,
            log_ex,
        )


# ======================================================================================
# HELPERS DE NORMALIZAÇÃO
# ======================================================================================

def _normalize_str(value: Any, max_length: int = 4000) -> Optional[str]:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    return text[:max_length]


def _normalize_ws(value: Any) -> Optional[str]:
    text = _normalize_str(value, max_length=100)
    if not text:
        return None

    normalized = text.strip().upper()
    if not re.fullmatch(r"WS-\d+", normalized):
        return None

    return normalized


def _normalize_compact(value: Any) -> Optional[str]:
    text = _normalize_str(value, max_length=200)
    if text is None:
        return None
    return re.sub(r"[\s\-_]+", "", text).lower()


def _normalize_stage_eligible(value: Any) -> Optional[str]:
    normalized = _normalize_compact(value)

    if normalized == "eligible":
        return "Y"

    if normalized == "noteligible":
        return "N"

    if normalized == "expired":
        return "Y"

    return None


def _normalize_opt_in_flag(value: Any) -> Optional[int]:
    normalized = _normalize_compact(value)

    if normalized == "optedin":
        return 1

    if normalized == "pending":
        return 0

    if normalized == "optedout":
        return 1

    return None


def _normalize_opt_in_label(value: Any) -> Optional[str]:
    normalized = _normalize_compact(value)

    if normalized == "optedin":
        return "OPTED_IN"

    if normalized == "pending":
        return "PENDING"

    if normalized == "optedout":
        return "OPTED_OUT"

    return None


def _parse_yes_no_flag(value: Any) -> int:
    normalized = _normalize_compact(value)

    if normalized in {"y", "yes", "true", "1"}:
        return 1

    if normalized in {"n", "no", "false", "0"}:
        return 0

    return 0


def _to_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()[:50]

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d %b %Y",
        "%d %B %Y",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def _to_decimal(value: Any) -> Optional[Decimal]:
    if value is None:
        return None

    if isinstance(value, Decimal):
        return value

    text = str(value).strip()
    if not text:
        return None

    text = re.sub(r"[^\d,\.\-+]", "", text)

    if not text or text in {"", ".", "-", "+", "-.", "+."}:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "")
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "")
        text = text.replace(",", ".")

    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _safe_int(value: Any, default: int = 0) -> int:
    if value is None:
        return default

    if isinstance(value, bool):
        return default

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        return int(value) if value.is_integer() else default

    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]

    if text.lstrip("+-").isdigit():
        try:
            return int(text)
        except Exception:
            return default

    return default


def _sanitize_decimal_30_6(value: Any) -> Optional[Decimal]:
    dec = _to_decimal(value)
    if dec is None:
        return None

    try:
        dec = dec.quantize(DECIMAL_30_6_QUANTIZER, rounding=ROUND_HALF_UP)
    except Exception:
        return None

    if dec < DECIMAL_30_6_MIN or dec > DECIMAL_30_6_MAX:
        return None

    return dec


def _normalize_compare_string(value: Any) -> Optional[str]:
    return _normalize_str(value, max_length=4000)


def _normalize_compare_decimal(value: Any) -> Optional[Decimal]:
    dec = _to_decimal(value)
    if dec is None:
        return None

    try:
        return dec.normalize()
    except Exception:
        return dec


def _normalize_compare_date(value: Any) -> Optional[date]:
    return _to_date(value)


def _values_different(field: str, old_val: Any, new_val: Any) -> bool:
    int_fields = {
        "task_status",
        "task_customer_id",
        "task_cr_party_id",
        "task_opt_in_flag",
        "task_ea_flag",
        "task_owner_id",
    }

    decimal_fields = {
        "task_booking_amount",
    }

    date_fields = {
        "task_start",
        "task_start_performed",
        "task_end",
        "task_end_performed",
        "task_booking_date",
    }

    string_fields = {
        "task_reference",
        "task_cr_party_name",
        "task_currency",
        "task_track",
        "task_subtrack",
        "task_architecture",
        "task_deal_id",
        "task_ws",
        "task_eligible",
        "task_status_justification",
    }

    if old_val is None and new_val is None:
        return False

    if field in int_fields:
        return _safe_int(old_val, default=0) != _safe_int(new_val, default=0)

    if field in decimal_fields:
        return _normalize_compare_decimal(old_val) != _normalize_compare_decimal(new_val)

    if field in date_fields:
        return _normalize_compare_date(old_val) != _normalize_compare_date(new_val)

    if field in string_fields:
        return _normalize_compare_string(old_val) != _normalize_compare_string(new_val)

    return old_val != new_val


# ======================================================================================
# HELPERS DE ARQUIVO EXCEL
# ======================================================================================

def _resolve_input_path(file_path: str) -> Path:
    path = Path(file_path)
    if path.is_absolute():
        return path
    return BASE_INPUT_PATH / path.name


def _build_failed_output_path(input_path: Path) -> Path:
    return BASE_OUTPUT_PATH / f"{input_path.stem}_failed_rows.xlsx"


def _open_workbook_rw(file_path: str) -> Tuple[Any, Worksheet, Dict[str, int], List[str]]:
    wb = load_workbook(filename=file_path, read_only=False, data_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    idx_map = {h: i for i, h in enumerate(headers)}
    return wb, ws, idx_map, headers


def _ensure_failed_workbook(
    failed_path: Path,
    original_headers: List[str],
    execution_log_path: Optional[Path] = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "failed_rows"
    ws.append(original_headers + ERROR_EXTRA_COLUMNS)
    wb.save(str(failed_path))
    wb.close()

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO failed file created (overwritten if existed): {str(failed_path)}",
        )


def _append_failed_rows(
    failed_path: Path,
    rows_to_append: List[List[Any]],
    original_headers: List[str],
    execution_log_path: Optional[Path] = None,
) -> None:
    if not rows_to_append:
        return

    _ensure_failed_workbook(failed_path, original_headers, execution_log_path)

    wb = load_workbook(str(failed_path))
    ws = wb.active

    for row in rows_to_append:
        ws.append(row)

    wb.save(str(failed_path))
    wb.close()

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO appended failed rows count={len(rows_to_append)} to file={str(failed_path)}",
        )


def _row_dict_to_failed_excel_row(
    headers: List[str],
    row_dict: Dict[str, Any],
    original_row_number: int,
    input_file_path: str,
    result: RowProcessResult,
) -> List[Any]:
    processed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    original_values = [row_dict.get(h) for h in headers]

    return original_values + [
        result.error_message,
        result.error_column,
        str(result.error_value)[:2000] if result.error_value is not None else None,
        original_row_number,
        processed_at,
        IMPORT_SOURCE,
        input_file_path,
    ]


def _get_row_value(row_dict: Dict[str, Any], logical_header: str) -> Any:
    aliases = HEADER_ALIASES.get(logical_header, [logical_header])
    for alias in aliases:
        if alias in row_dict:
            return row_dict.get(alias)
    return None


# ======================================================================================
# HELPERS DE TASK / ACTIVITY / USE CASE
# ======================================================================================

def _validate_required_columns(headers: List[str]) -> List[str]:
    missing: List[str] = []

    for col in REQUIRED_COLUMNS:
        aliases = HEADER_ALIASES.get(col, [col])
        if not any(alias in headers for alias in aliases):
            missing.append(col)

    return missing


def _normalize_stage_detail_text(raw_value: Any) -> Optional[str]:
    text = _normalize_str(raw_value, max_length=12000)
    if not text:
        return None

    normalized = text.replace("\r", "\n")
    normalized = re.sub(r"\u00a0", " ", normalized)
    normalized = re.sub(r"\s+\n", "\n", normalized)
    normalized = re.sub(r"\n\s+", "\n", normalized)
    normalized = re.sub(r"\s{2,}", " ", normalized).strip()

    if not normalized:
        return None

    if "\n" in normalized:
        lines = [line.strip() for line in normalized.split("\n") if line.strip()]
        return "\n".join(lines) if lines else None

    split_text = re.sub(r"\s+(?=\d+\.\s)", "\n", normalized)
    lines = [line.strip() for line in split_text.split("\n") if line.strip()]
    return "\n".join(lines) if lines else None


def _remove_yes_no_suffix(text: str) -> str:
    return re.sub(r"\s*-\s*[YyNn]\s*$", "", text).strip()


def _parse_stage_details(raw_value: Any) -> Dict[str, Any]:
    normalized_text = _normalize_stage_detail_text(raw_value)

    if not normalized_text:
        return {
            "raw_scope": None,
            "clean_scope": None,
            "criteria_total": 0,
            "criteria_yes": 0,
            "activity_completed": Decimal("0.000000"),
        }

    lines = [line.strip() for line in normalized_text.split("\n") if line.strip()]
    criteria_yes = 0
    clean_lines: List[str] = []

    for line in lines:
        if re.search(r"\s*-\s*[Yy]\s*$", line):
            criteria_yes += 1
        clean_lines.append(_remove_yes_no_suffix(line))

    criteria_total = len(lines)
    denominator = criteria_total + 1
    completed = Decimal("0.000000")
    if denominator > 0:
        completed = (
            Decimal(criteria_yes) / Decimal(denominator)
        ).quantize(DECIMAL_30_6_QUANTIZER, rounding=ROUND_HALF_UP)

    return {
        "raw_scope": "\n".join(lines),
        "clean_scope": "\n".join(clean_lines) if clean_lines else None,
        "criteria_total": criteria_total,
        "criteria_yes": criteria_yes,
        "activity_completed": completed,
    }


def _calc_activity_completed_for_stage(
    stage_name: str,
    criteria_yes: int,
    criteria_total: int,
) -> Decimal:
    """
    Calcula activity_completed conforme regra de negócio por estágio.

    Regra geral  (Onboard, Use, Engage, Adopt):
        activity_completed = criteria_yes / (criteria_total + 1)

    Exceção Implement:
        activity_completed = criteria_yes / criteria_total
        (sem o "+ 1" no denominador)
    """
    if stage_name == "Implement":
        if criteria_total <= 0:
            return Decimal("0.000000")
        return (
            Decimal(criteria_yes) / Decimal(criteria_total)
        ).quantize(DECIMAL_30_6_QUANTIZER, rounding=ROUND_HALF_UP)

    denominator = criteria_total + 1
    if denominator <= 0:
        return Decimal("0.000000")
    return (
        Decimal(criteria_yes) / Decimal(denominator)
    ).quantize(DECIMAL_30_6_QUANTIZER, rounding=ROUND_HALF_UP)


def _normalize_row(row_dict: Dict[str, Any]) -> Dict[str, Any]:
    stage_value = _get_row_value(row_dict, "Stage")
    opt_in_value = _get_row_value(row_dict, "Adopt Rebate Opt-In Status")
    incentive_level = _get_row_value(row_dict, "Incentive Level")

    stage_code = _normalize_stage_eligible(stage_value)
    opt_in_flag = _normalize_opt_in_flag(opt_in_value)
    opt_in_label = _normalize_opt_in_label(opt_in_value)

    stage_details: Dict[str, Dict[str, Any]] = {}
    stage_dates: Dict[str, Optional[date]] = {}
    stage_flags: Dict[str, int] = {}

    for stage_name, config in STAGE_CONFIG.items():
        stage_details[stage_name] = _parse_stage_details(_get_row_value(row_dict, config["details_header"]))
        stage_dates[stage_name] = _to_date(_get_row_value(row_dict, config["date_header"]))
        stage_flags[stage_name] = _parse_yes_no_flag(_get_row_value(row_dict, config["flag_header"]))

    normalized = {
        "ws": _normalize_ws(_get_row_value(row_dict, "Deal WS-ID")),
        "cr_party_name": _normalize_str(_get_row_value(row_dict, "CR Party Name")),
        "cr_party_id": _safe_int(_get_row_value(row_dict, "CR Party ID"), default=0),
        "track": _normalize_str(_get_row_value(row_dict, "Track")),
        "subtrack": _normalize_str(_get_row_value(row_dict, "Sub-Track")),
        "incentive_level": _normalize_str(incentive_level),
        "task_reference": _normalize_str(incentive_level),
        "task_end": _to_date(_get_row_value(row_dict, "Deal Incentive Expiry Date")),
        "task_start": _to_date(_get_row_value(row_dict, "Booking Date")),
        "task_booking_date": _to_date(_get_row_value(row_dict, "Booking Date")),
        "task_booking_amount": _sanitize_decimal_30_6(_get_row_value(row_dict, "Booking Amount - Net to Cisco")),
        "task_eligible": stage_code,
        "stage_raw": _normalize_str(stage_value),
        "task_opt_in_flag": opt_in_flag,
        "opt_in_raw": _normalize_str(opt_in_value),
        "opt_in_label": opt_in_label,
        "task_start_performed": _to_date(_get_row_value(row_dict, "Adopt Rebate Start Date")),
        "task_ea_flag": _parse_yes_no_flag(_get_row_value(row_dict, "EA Flag")),
        "task_deal_id": _normalize_str(_get_row_value(row_dict, "Deal ID")),
        "task_architecture": _normalize_str(_get_row_value(row_dict, "Deal CPI Portfolio")),
        "subscription_id": _normalize_str(_get_row_value(row_dict, "Subscription ID")),
        "booking_po_number": _normalize_str(_get_row_value(row_dict, "Booking PO Number")),
        "task_currency": "USD",
        "stage_details": stage_details,
        "stage_dates": stage_dates,
        "stage_flags": stage_flags,
    }

    if normalized["task_start"] and normalized["task_start_performed"]:
        if normalized["task_start"] < normalized["task_start_performed"]:
            normalized["task_start"] = normalized["task_start_performed"]

    return normalized


def _validate_minimum_required_fields(data: Dict[str, Any]) -> Tuple[bool, Optional[str], Optional[Any]]:
    required_map = {
        "Deal WS-ID": data.get("ws"),
        "CR Party Name": data.get("cr_party_name"),
        "CR Party ID": data.get("cr_party_id"),
        "Track": data.get("track"),
        "Sub-Track": data.get("subtrack"),
        "Deal Incentive Expiry Date": data.get("task_end"),
        "Booking Date": data.get("task_booking_date"),
        "Stage": data.get("task_eligible"),
        "Adopt Rebate Opt-In Status": data.get("task_opt_in_flag"),
        "Deal ID": data.get("task_deal_id"),
        "Deal CPI Portfolio": data.get("task_architecture"),
    }

    for col, value in required_map.items():
        if value is None or value == "":
            return False, col, value

    if data.get("cr_party_id", 0) <= 0:
        return False, "CR Party ID", data.get("cr_party_id")

    requires_adopt_rebate_start_date = (
        data.get("task_eligible") == "Y"
        and data.get("opt_in_label") == "OPTED_IN"
    )

    if requires_adopt_rebate_start_date and data.get("task_start_performed") is None:
        return False, "Adopt Rebate Start Date", data.get("task_start_performed")

    return True, None, None


def _resolve_company_id(company_name: Optional[str]) -> int:
    if not company_name:
        return 0

    try:
        company_id = repo_company.get_company_id_by_name(company_name)
        return _safe_int(company_id, default=0)
    except Exception:
        return 0


def _resolve_existing_task(data: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    ws = data.get("ws")
    if not ws:
        return None

    try:
        found_ids = repo_task.find_ids_by({"task_ws": ws})
    except Exception:
        return None

    if not found_ids:
        return None

    matched_tasks: List[Dict[str, Any]] = []

    for found_id in found_ids:
        task_id = _safe_int(found_id, default=0)
        if not task_id:
            continue

        task_row = _get_task_columns_for_match(task_id)
        if not task_row:
            continue

        task_type_id = _safe_int(task_row.get("task_tasktype_id"), default=0)
        if task_type_id in VALID_EXISTING_TASK_TYPE_IDS:
            matched_tasks.append(task_row)

    if not matched_tasks:
        return None

    type_22_tasks = [
        t for t in matched_tasks
        if _safe_int(t.get("task_tasktype_id"), default=0) == TASK_TYPE_LCI_CURRENT
    ]
    type_21_tasks = [
        t for t in matched_tasks
        if _safe_int(t.get("task_tasktype_id"), default=0) == TASK_TYPE_LCI_LEGACY
    ]

    if len(type_22_tasks) == 1:
        return type_22_tasks[0]
    if len(type_22_tasks) > 1:
        raise ValueError(
            f"Mais de uma task tipo 22 encontrada para o mesmo WS {ws}. "
            f"task_ids={[t.get('task_id') for t in type_22_tasks]}"
        )

    if len(type_21_tasks) == 1:
        return type_21_tasks[0]
    if len(type_21_tasks) > 1:
        raise ValueError(
            f"Mais de uma task tipo 21 encontrada para o mesmo WS {ws}. "
            f"task_ids={[t.get('task_id') for t in type_21_tasks]}"
        )

    raise ValueError(
        f"Mais de uma task encontrada para o mesmo WS {ws}. "
        f"task_ids={[t.get('task_id') for t in matched_tasks]}"
    )


def _get_task_columns_for_match(task_id: int) -> Optional[Dict[str, Any]]:
    required_columns = [
        "task_id",
        "task_tasktype_id",
        "task_status",
        "task_owner_id",
        "task_opt_in_flag",
        "task_ea_flag",
        "task_reference",
        "task_customer_id",
        "task_cr_party_id",
        "task_cr_party_name",
        "task_end",
        "task_currency",
        "task_track",
        "task_subtrack",
        "task_booking_date",
        "task_booking_amount",
        "task_architecture",
        "task_deal_id",
        "task_ws",
        "task_start",
        "task_start_performed",
        "task_end_performed",
        "task_eligible",
        "task_status_justification",
        "task_completed",
    ]

    try:
        row = repo_task.get_columns_by_task_id(task_id, required_columns, as_df=False)
        if isinstance(row, dict):
            return row
        if isinstance(row, list) and row and isinstance(row[0], dict):
            return row[0]
    except Exception:
        return None

    return None


def _build_task_payload(
    data: Dict[str, Any],
    company_id: int,
    task_type_id: int,
) -> Dict[str, Any]:
    is_opted_out = data.get("opt_in_label") == "OPTED_OUT"

    return {
        "task_tasktype_id": task_type_id,
        "task_reference": data.get("task_reference"),
        "task_owner_id": DEFAULT_TASK_OWNER_ID,
        "task_temp_owner_id": None,
        "task_customer_id": company_id,
        "task_cr_party_id": data.get("cr_party_id"),
        "task_cr_party_name": data.get("cr_party_name"),
        "task_created_in": datetime.now(),
        "task_created_by": 0,
        "task_priority": "LOW",
        "task_project_id": 0,
        "task_status": (
            TASK_STATUS_CANCELLED
            if is_opted_out
            else (TASK_STATUS_IN_PROGRESS if data.get("task_opt_in_flag") == 1 else TASK_STATUS_OPEN)
        ),
        "task_status_justification": None,
        "task_start": data.get("task_start"),
        "task_end": data.get("task_end"),
        "task_start_performed": data.get("task_start_performed"),
        "task_end_performed": data.get("task_end"),
        "task_value": Decimal("0.000000"),
        "task_forecast": 0,
        "task_backlog": 0,
        "task_rate": 1,
        "task_currency": "USD",
        "task_ws": data.get("ws"),
        "task_deal_id": data.get("task_deal_id"),
        "task_track": data.get("track"),
        "task_subtrack": data.get("subtrack"),
        "task_highlight": 0,
        "task_remark": None,
        "task_description": None,
        "task_ea_flag": data.get("task_ea_flag", 0),
        "task_opt_in_flag": data.get("task_opt_in_flag", 0),
        "task_completed": 0,
        "task_architecture": data.get("task_architecture"),
        "task_solution_domain": None,
        "task_eligible": "N" if is_opted_out else data.get("task_eligible"),
        "task_end_fy": None,
        "task_booking_date": data.get("task_booking_date"),
        "task_booking_amount": data.get("task_booking_amount"),
    }


def _insert_task_history(
    task_id: int,
    remark: str,
    updated_by: str = "System BA",
    activity_id: int = 0,
) -> None:
    history = {
        "taskrecord_task_id": task_id,
        "taskrecord_activity_id": activity_id,
        "taskrecord_remark": remark,
        "taskrecord_updated_by": updated_by,
        "taskrecord_type": "LOG",
    }
    repo_history.insert(history)


def _insert_opt_in_status_history(
    task_id: int,
    opt_in_raw: str,
    record_date: Any,
    updated_by: str = "System BA",
) -> None:
    """Grava um registro de tipo 'OPT IN STATUS' em tbTaskRecord.

    taskrecord_remark = valor original da coluna 'Adopt Rebate Opt-In Status' do Excel.
    taskrecord_date   = Adopt Rebate Start Date se não nulo, senão Booking Date.
    """
    history = {
        "taskrecord_task_id": task_id,
        "taskrecord_activity_id": 0,
        "taskrecord_date": record_date,
        "taskrecord_remark": opt_in_raw,
        "taskrecord_type": "OPT IN STATUS",
        "taskrecord_updated_by": updated_by,
    }
    repo_history.insert(history)


def _sync_opt_in_status_history(
    task_id: int,
    opt_in_raw: Optional[str],
    record_date: Any,
    is_new_task: bool = False,
    updated_by: str = "System BA",
) -> None:
    """
    Registra histórico de mudança de 'Adopt Rebate Opt-In Status' em tbTaskRecord.

    opt_in_raw   = valor original da coluna 'Adopt Rebate Opt-In Status' do Excel.
    record_date  = Adopt Rebate Start Date se não nulo, senão Booking Date.

    - Para task nova: sempre insere.
    - Para task existente: insere somente se o último registro 'OPT IN STATUS'
      tiver remark diferente do valor atual lido do Excel.
    """
    if not opt_in_raw:
        return

    if is_new_task:
        _insert_opt_in_status_history(
            task_id=task_id,
            opt_in_raw=opt_in_raw,
            record_date=record_date,
            updated_by=updated_by,
        )
        return

    try:
        last_record = repo_history.get_last_opt_in_status(task_id=task_id)
    except Exception:
        last_record = None

    last_remark = (
        _normalize_str(last_record.get("taskrecord_remark"))
        if last_record
        else None
    )

    if last_remark != opt_in_raw:
        _insert_opt_in_status_history(
            task_id=task_id,
            opt_in_raw=opt_in_raw,
            record_date=record_date,
            updated_by=updated_by,
        )


def _activity_status_name(status_id: int) -> str:
    status_map = {
        1: "OPEN",
        2: "IN PROGRESS",
        3: "ON HOLD",
        4: "CANCELLED",
        5: "DECLINED",
        6: "EXPIRED",
        7: "SUBMITTED",
        8: "RESUBMITTED",
        9: "APPROVED TO CLOSE",
        10: "COMPLETED",
    }
    return status_map.get(status_id, f"STATUS {status_id}")


def _insert_activity_status_change_history(
    task_id: int,
    activity_id: int,
    new_status: int,
    updated_by: str = "System BA",
) -> None:
    history = {
        "taskrecord_task_id": task_id,
        "taskrecord_activity_id": activity_id,
        "taskrecord_date": datetime.now(),
        "taskrecord_remark": _activity_status_name(new_status),
        "taskrecord_type": "STATUS CHANGE",
        "taskrecord_updated_by": updated_by,
    }
    repo_history.insert(history)


def _confirm_created_task(task_id: int, expected_ws: Optional[str]) -> bool:
    if not task_id:
        return False

    task_row = _get_task_columns_for_match(task_id)
    if not task_row:
        return False

    confirmed_task_id = _safe_int(task_row.get("task_id"), default=0)
    if confirmed_task_id != task_id:
        return False

    if expected_ws is not None:
        confirmed_ws = _normalize_ws(task_row.get("task_ws"))
        if confirmed_ws != expected_ws:
            return False

    return True


def _close_task_as_not_eligible(task_id: int, updated_by: str = "System BA") -> int:
    try:
        repo_task.update(
            data={
                "task_status": TASK_STATUS_DECLINED,
                "task_eligible": "N",
                "task_status_justification": TASK_NOT_ELIGIBLE_JUSTIFICATION,
            },
            where={"task_id": task_id},
        )

        _insert_task_history(
            task_id=task_id,
            activity_id=0,
            remark="Task status changed to Declined; Task marked as not eligible by Cisco",
            updated_by=updated_by,
        )

        return 0
    except Exception:
        return 0


def _close_open_activities_as_cancelled(task_id: int, updated_by: str = "System BA") -> int:
    try:
        activities = repo_activity.get_activity(task_id=task_id, as_df=False)
    except Exception:
        return 0

    cancelled_count = 0

    for activity in activities or []:
        activity_id = _safe_int(activity.get("activity_id"), default=0)
        activity_status = _safe_int(activity.get("activity_status"), default=0)

        if not activity_id or activity_status in CLOSED_ACTIVITY_STATUSES:
            continue

        try:
            rows = repo_activity.update(
                data={"activity_status": ACTIVITY_STATUS_CANCELLED},
                where={"activity_id": activity_id},
            )
        except Exception:
            rows = 0

        if rows > 0:
            cancelled_count += 1
            _insert_activity_status_change_history(
                task_id=task_id,
                activity_id=activity_id,
                new_status=ACTIVITY_STATUS_CANCELLED,
                updated_by=updated_by,
            )
            _insert_task_history(
                task_id=task_id,
                activity_id=activity_id,
                remark="Opt-Out was performed for this task",
                updated_by=updated_by,
            )

    return cancelled_count


def _apply_opted_out_rule_existing_task(
    task_id: int,
    current_status: int,
    updated_by: str = "System BA",
) -> Tuple[bool, int]:
    if current_status in CLOSED_TASK_STATUSES:
        return False, 0

    repo_task.update(
        data={
            "task_status": TASK_STATUS_CANCELLED,
            "task_forecast": 0,
            "task_eligible": "N",
        },
        where={"task_id": task_id},
    )

    _insert_task_history(
        task_id=task_id,
        activity_id=0,
        remark="Opt-Out was performed for this task",
        updated_by=updated_by,
    )

    activities_cancelled = _close_open_activities_as_cancelled(
        task_id=task_id,
        updated_by=updated_by,
    )

    return True, activities_cancelled


def _reopen_task_if_needed(task_id: int, updated_by: str = "System BA") -> bool:
    try:
        rows = repo_task.update(
            data={
                "task_status": TASK_STATUS_IN_PROGRESS,
                "task_status_justification": None,
                "task_eligible": "Y",
            },
            where={"task_id": task_id},
        )
        if rows > 0:
            _insert_task_history(
                task_id=task_id,
                activity_id=0,
                remark="Task reopened as In Progress due to eligible Opted In source row",
                updated_by=updated_by,
            )
            return True
    except Exception:
        return False

    return False


def _find_matching_activity(task_id: int, stage_name: str) -> Optional[Dict[str, Any]]:
    try:
        activities = repo_activity.get_activity(task_id=task_id, as_df=False)
    except Exception:
        return None

    normalized_target = _normalize_compact(stage_name)

    for activity in activities:
        if _normalize_compact(activity.get("activity_name")) == normalized_target:
            return activity

    return None


def _create_missing_implement_activity(
    task_id: int,
    data: Dict[str, Any],
    file_path: str,
    row_number: int,
    updated_by: str = "System BA",
) -> bool:
    """
    Cria a activity 'Implement' quando a task já possui filhas mas a atividade
    Implement ainda não existe e a coluna 'Task Details (Implement)' não está vazia.

    Regra exclusiva para Implement.

    Parâmetros:
        task_id    : ID da task pai
        data       : dicionário normalizado da linha Excel
        file_path  : caminho do arquivo de origem (para log)
        row_number : número da linha lida (para log)
        updated_by : responsável pelo registro de histórico

    Retorna:
        True  → activity criada com sucesso
        False → falha na criação (erro já logado)
    """
    stage_payload = data["stage_details"].get("Implement", {})
    raw_scope = stage_payload.get("raw_scope")

    if not raw_scope:
        return False

    stage_flag = data["stage_flags"].get("Implement", 0)
    stage_date = data["stage_dates"].get("Implement")

    # Determinar activity_status e activity_completed
    if stage_flag == 1 and stage_date is not None:
        activity_status = ACTIVITY_STATUS_COMPLETED
        activity_completed = Decimal("1.000000")
    else:
        activity_status = ACTIVITY_STATUS_IN_PROGRESS
        # Implement usa formula criteria_yes / criteria_total (sem +1)
        criteria_total = stage_payload.get("criteria_total", 0)
        criteria_yes = stage_payload.get("criteria_yes", 0)
        activity_completed = _calc_activity_completed_for_stage(
            stage_name="Implement",
            criteria_yes=criteria_yes,
            criteria_total=criteria_total,
        )

    # activity_end_performed: usar Stage Completion Date (Implement); se nulo, usar Deal Incentive Expiry Date
    activity_end_performed = stage_date if stage_date is not None else data.get("task_end")

    # activity_start e activity_start_performed: usar Adopt Rebate Start Date
    activity_start = data.get("task_start_performed")
    activity_start_performed = data.get("task_start_performed")

    insert_data: Dict[str, Any] = {
        "activity_task_id": task_id,
        "activity_name": "Implement",
        "activity_seq": STAGE_CONFIG["Implement"]["seq"],
        "activity_status": activity_status,
        "activity_scope": raw_scope,
        "activity_completed": activity_completed,
        "activity_deal_id": data.get("task_deal_id"),
        "activity_track": data.get("track"),
        "activity_sub_track": data.get("subtrack"),
        "activity_value": Decimal("0.000000"),
        "activity_currency": "USD",
        "activity_start": activity_start,
        "activity_end": data.get("task_end"),
        "activity_start_performed": activity_start_performed,
        "activity_end_performed": activity_end_performed,
        "activity_approved": 0,
        "activity_backlog_value": Decimal("0.000000"),
    }

    try:
        new_activity_id = repo_activity.insert(insert_data)
    except Exception as ex:
        _safe_log(
            file_path,
            row_number,
            f"Falha ao criar activity 'Implement' ausente: {ex}",
            "Task Details (Implement)",
            raw_scope[:200] if raw_scope else None,
        )
        return False

    if not new_activity_id:
        _safe_log(
            file_path,
            row_number,
            "Falha ao criar activity 'Implement' ausente: insert retornou id=0",
            "Task Details (Implement)",
            raw_scope[:200] if raw_scope else None,
        )
        return False

    _insert_activity_status_change_history(
        task_id=task_id,
        activity_id=new_activity_id,
        new_status=activity_status,
        updated_by=updated_by,
    )

    _insert_task_history(
        task_id=task_id,
        activity_id=new_activity_id,
        remark=(
            f"Activity 'Implement' criada via {IMPORT_SOURCE} "
            f"(estágio ausente detectado para task com outras activities existentes); "
            f"activity_status={activity_status}; activity_completed={activity_completed}"
        ),
        updated_by=updated_by,
    )

    return True


def _update_activity_from_stage(
    task_id: int,
    stage_name: str,
    stage_payload: Dict[str, Any],
    updated_by: str = "System BA",
) -> Tuple[bool, bool]:
    activity = _find_matching_activity(task_id, stage_name)
    if not activity:
        return False, False

    activity_id = _safe_int(activity.get("activity_id"), default=0)
    if not activity_id:
        return False, False

    current_status = _safe_int(activity.get("activity_status"), default=0)

    updates: Dict[str, Any] = {}
    remarks: List[str] = []

    new_scope = stage_payload.get("raw_scope")
    if new_scope is not None and _normalize_compare_string(activity.get("activity_scope")) != _normalize_compare_string(new_scope):
        updates["activity_scope"] = new_scope
        remarks.append(f"Change activity_scope for {stage_name}")

    allow_completed_update = current_status not in {ACTIVITY_STATUS_COMPLETED}
    if current_status in {
        ACTIVITY_STATUS_CANCELLED,
        ACTIVITY_STATUS_DECLINED,
        ACTIVITY_STATUS_EXPIRED,
    }:
        allow_completed_update = True

    if stage_name == "Implement":
        criteria_total = stage_payload.get("criteria_total", 0)
        criteria_yes = stage_payload.get("criteria_yes", 0)
        new_completed = _calc_activity_completed_for_stage(
            stage_name="Implement",
            criteria_yes=criteria_yes,
            criteria_total=criteria_total,
        )
    else:
        new_completed = stage_payload.get("activity_completed")

    if (
        allow_completed_update
        and new_completed is not None
        and _normalize_compare_decimal(activity.get("activity_completed")) != _normalize_compare_decimal(new_completed)
    ):
        updates["activity_completed"] = new_completed
        remarks.append(f"Change activity_completed for {stage_name} to {new_completed}")

    if updates:
        old_status = current_status
        new_status = _safe_int(updates.get("activity_status"), default=old_status)

        repo_activity.update(data=updates, where={"activity_id": activity_id})

        if "activity_status" in updates and new_status != old_status:
            _insert_activity_status_change_history(
                task_id=task_id,
                activity_id=activity_id,
                new_status=new_status,
                updated_by=updated_by,
            )

        if remarks:
            _insert_task_history(
                task_id=task_id,
                activity_id=activity_id,
                remark="; ".join(remarks),
                updated_by=updated_by,
            )
        return True, True

    return True, False


def _sync_task_activities(
    task_id: int,
    data: Dict[str, Any],
    file_path: str,
    row_number: int,
) -> Tuple[bool, int, Optional[str], Optional[str], Optional[Any], int]:
    activity_count = repo_activity.get_count_activity(task_id)

    if activity_count == 0:
        if data.get("task_eligible") == "Y" and data.get("task_opt_in_flag") != 0:
            return (
                False,
                0,
                "Task elegível com Opted In, porém sem activities cadastradas",
                "Deal WS-ID",
                data.get("ws"),
                0,
            )
        return True, 0, None, None, None, 0

    # ------------------------------------------------------------------
    # REGRA EXCLUSIVA IMPLEMENT: criar activity se ausente e detalhes
    # disponíveis no Excel
    # ------------------------------------------------------------------
    if (
        _find_matching_activity(task_id, "Implement") is None
        and data["stage_details"].get("Implement", {}).get("raw_scope")
    ):
        implement_created = _create_missing_implement_activity(
            task_id=task_id,
            data=data,
            file_path=file_path,
            row_number=row_number,
        )
        if implement_created:
            # Re-conta activities após a criação para manter
            # activity_count preciso ao final da função
            activity_count = repo_activity.get_count_activity(task_id)
            _safe_log(
                file_path,
                row_number,
                "Activity 'Implement' criada automaticamente (estágio ausente detectado)",
                "Task Details (Implement)",
                data.get("ws"),
            )

    updated_count = 0

    for stage_name, config in STAGE_CONFIG.items():
        stage_payload = data["stage_details"].get(stage_name, {})

        found, updated = _update_activity_from_stage(
            task_id=task_id,
            stage_name=stage_name,
            stage_payload=stage_payload,
        )

        if not found:
            _safe_log(
                file_path,
                row_number,
                f"Activity '{stage_name}' não encontrada para a task",
                "Deal WS-ID",
                data.get("ws"),
            )
            continue

        if updated:
            updated_count += 1

    return True, updated_count, None, None, None, activity_count


def _ensure_use_case(data: Dict[str, Any]) -> int:
    if repo_use_case is None:
        return 0

    try:
        uc_ids = repo_use_case.get_use_case_ids(
            vendor_id=VENDOR_ID_CISCO,
            track=data.get("track"),
            subtrack=data.get("subtrack"),
        )
    except Exception:
        uc_ids = []

    if uc_ids:
        return _safe_int(uc_ids[0], default=0)

    try:
        return _safe_int(
            repo_use_case.insert_use_case(
                {
                    "uc_vendor_id": VENDOR_ID_CISCO,
                    "uc_architecture": data.get("task_architecture"),
                    "uc_track": data.get("track"),
                    "uc_use_case": data.get("subtrack"),
                    "uc_update_date": datetime.now(),
                }
            ),
            default=0,
        )
    except Exception:
        return 0


def _should_update_exit_criteria(update_date: Optional[datetime]) -> bool:
    if update_date is None:
        return True

    today = datetime.now().date()
    return update_date.date() < today


def _sync_use_case_exit_criteria(data: Dict[str, Any], uc_id: int) -> None:
    if repo_use_case is None or not uc_id:
        return

    for stage_name, config in STAGE_CONFIG.items():
        scope_clean = data["stage_details"][stage_name].get("clean_scope")
        if not scope_clean:
            continue

        try:
            ucec_ids = repo_use_case.get_use_case_exit_criteria_ids(
                vendor_id=VENDOR_ID_CISCO,
                track=data.get("track"),
                subtrack=data.get("subtrack"),
                name=stage_name,
            )
        except Exception:
            ucec_ids = []

        ucec_id = _safe_int(ucec_ids[0], default=0) if ucec_ids else 0

        if ucec_id:
            try:
                update_date = repo_use_case.get_use_case_exit_criteria_update_date(ucec_id=ucec_id)
            except Exception:
                update_date = None

            if _should_update_exit_criteria(update_date):
                try:
                    repo_use_case.update_exit_criteria(
                        ucec_id=ucec_id,
                        data={
                            "ucec_scope": scope_clean,
                            "ucec_update_date": datetime.now().date(),
                            "ucec_updated_by": "System BA",
                        },
                    )
                except Exception:
                    continue
            continue

        try:
            repo_use_case.insert_exit_criteria(
                {
                    "ucec_tasktype_id": TASK_TYPE_LCI_CURRENT,
                    "ucec_uc_id": uc_id,
                    "ucec_seq": config["seq"],
                    "ucec_name": stage_name,
                    "ucec_scope": scope_clean,
                    "ucec_update_date": datetime.now().date(),
                    "ucec_updated_by": "System BA",
                }
            )
        except Exception:
            continue


def _update_existing_task(
    task_id: int,
    existing_task: Dict[str, Any],
    payload: Dict[str, Any],
    updated_by: str = "System BA",
) -> bool:
    updates: Dict[str, Any] = {}
    remarks: List[str] = []

    simple_compare_fields = [
        "task_reference",
        "task_customer_id",
        "task_cr_party_id",
        "task_cr_party_name",
        "task_end",
        "task_currency",
        "task_track",
        "task_subtrack",
        "task_booking_date",
        "task_booking_amount",
        "task_architecture",
        "task_deal_id",
        "task_ws",
        "task_eligible",
        "task_opt_in_flag",
        "task_ea_flag",
    ]

    for field in simple_compare_fields:
        old_val = existing_task.get(field)
        new_val = payload.get(field)

        if new_val is None:
            continue

        if _values_different(field, old_val, new_val):
            updates[field] = new_val
            remarks.append(f"Change {field} to {new_val}")

    payload_task_start = payload.get("task_start")
    if payload_task_start is not None and _values_different("task_start", existing_task.get("task_start"), payload_task_start):
        updates["task_start"] = payload_task_start
        remarks.append(f"Change task_start to {payload_task_start}")

    payload_start_performed = payload.get("task_start_performed")
    current_start_performed = _normalize_compare_date(existing_task.get("task_start_performed"))
    new_start_performed = _normalize_compare_date(payload_start_performed)
    if new_start_performed is not None:
        if current_start_performed is None or current_start_performed < new_start_performed:
            updates["task_start_performed"] = payload_start_performed
            remarks.append(f"Change task_start_performed to {payload_start_performed}")

    payload_end_performed = payload.get("task_end_performed")
    if payload_end_performed is not None and _values_different("task_end_performed", existing_task.get("task_end_performed"), payload_end_performed):
        updates["task_end_performed"] = payload_end_performed
        remarks.append(f"Change task_end_performed to {payload_end_performed}")

    if not updates:
        return False

    repo_task.update(data=updates, where={"task_id": task_id})

    _insert_task_history(
        task_id=task_id,
        remark="; ".join(remarks),
        updated_by=updated_by,
    )
    return True


def _recalculate_task_completed_if_needed(
    task_id: int,
    current_task_status: int,
    current_task_completed: Any,
    updated_by: str = "System BA",
) -> bool:
    if not task_id:
        return False

    if current_task_status in CLOSED_TASK_STATUSES:
        return False

    total_activities, avg_completed = repo_task.get_task_completion_summary(task_id)

    if total_activities <= 0:
        return False

    current_completed_decimal = _normalize_compare_decimal(current_task_completed)
    new_completed_decimal = _normalize_compare_decimal(avg_completed)

    if current_completed_decimal == new_completed_decimal:
        return False

    rows = repo_task.update(
        data={"task_completed": avg_completed},
        where={"task_id": task_id},
    )

    if rows > 0:
        _insert_task_history(
            task_id=task_id,
            activity_id=0,
            remark=f"Change task_completed to {avg_completed}",
            updated_by=updated_by,
        )
        return True

    return False


def _log_numeric_sanitization_if_needed(
    file_path: str,
    row_number: int,
    original_value: Any,
    sanitized_value: Optional[Decimal],
    column_name: str,
    field_name: str,
) -> None:
    if original_value is None or str(original_value).strip() == "":
        return

    if sanitized_value is not None:
        return

    _safe_log(
        file_path,
        row_number,
        f"Valor inválido ou fora do range permitido para {field_name} DECIMAL(30,6). Campo será gravado como NULL.",
        column_name,
        original_value,
    )


# ======================================================================================
# PROCESSAMENTO DE LINHA
# ======================================================================================

def _process_single_row(
    row_dict: Dict[str, Any],
    file_path: str,
    row_number: int,
    execution_log_path: Optional[Path] = None,
) -> RowProcessResult:
    try:
        data = _normalize_row(row_dict)

        _log_numeric_sanitization_if_needed(
            file_path=file_path,
            row_number=row_number,
            original_value=_get_row_value(row_dict, "Booking Amount - Net to Cisco"),
            sanitized_value=data.get("task_booking_amount"),
            column_name="Booking Amount - Net to Cisco",
            field_name="task_booking_amount",
        )

        valid, invalid_col, invalid_val = _validate_minimum_required_fields(data)
        if not valid:
            message = f"Campo obrigatório ausente ou inválido: {invalid_col}"
            _safe_log(file_path, row_number, message, invalid_col, invalid_val)
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column=invalid_col,
                error_value=invalid_val,
            )

        company_id = _resolve_company_id(data.get("cr_party_name"))
        if not company_id:
            message = "Customer não encontrado a partir de 'CR Party Name'"
            _safe_log(file_path, row_number, message, "CR Party Name", data.get("cr_party_name"))
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="CR Party Name",
                error_value=data.get("cr_party_name"),
            )

        existing_task = _resolve_existing_task(data)
        payload = _build_task_payload(
            data=data,
            company_id=company_id,
            task_type_id=PRIORITY_TASK_TYPE_ID,
        )

        if existing_task is None:
            task_id = _safe_int(repo_task.insert(payload), default=0)
            if not task_id:
                raise ValueError("Falha ao criar task")

            if not _confirm_created_task(task_id=task_id, expected_ws=payload.get("task_ws")):
                message = "Falha ao confirmar persistência da task após criação"
                _safe_log(file_path, row_number, message, "Deal WS-ID", payload.get("task_ws"))
                return RowProcessResult(
                    success=False,
                    error_message=message,
                    error_column="Deal WS-ID",
                    error_value=payload.get("task_ws"),
                )

            _insert_task_history(
                task_id=task_id,
                remark=f"Task criada via {IMPORT_SOURCE}",
            )

            if data.get("opt_in_label") == "OPTED_OUT":
                _insert_task_history(
                    task_id=task_id,
                    remark="Opt-Out was performed for this task",
                )

            _sync_opt_in_status_history(
                task_id=task_id,
                opt_in_raw=data.get("opt_in_raw"),
                record_date=data.get("task_start_performed") or data.get("task_booking_date"),
                is_new_task=True,
            )

            uc_id = _ensure_use_case(data)
            _sync_use_case_exit_criteria(data, uc_id)

            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO row={row_number} task created task_id={task_id} ws={payload.get('task_ws')}",
                )

            return RowProcessResult(
                success=True,
                created=True,
                cancelled=(data.get("opt_in_label") == "OPTED_OUT"),
            )

        task_id = _safe_int(existing_task.get("task_id"), default=0)
        if not task_id:
            raise ValueError("task_id inválido na task existente")

        current_status = _safe_int(existing_task.get("task_status"), default=0)
        current_eligible = _normalize_str(existing_task.get("task_eligible"), max_length=10)
        opt_in_label = data.get("opt_in_label")

        if opt_in_label == "OPTED_OUT":
            _sync_opt_in_status_history(
                task_id=task_id,
                opt_in_raw=data.get("opt_in_raw"),
                record_date=data.get("task_start_performed") or data.get("task_booking_date"),
                is_new_task=False,
            )

            if current_status in CLOSED_TASK_STATUSES:
                if execution_log_path:
                    _append_execution_log(
                        execution_log_path,
                        f"INFO row={row_number} ignored opted-out already closed task_id={task_id} ws={data.get('ws')}",
                    )
                return RowProcessResult(success=True, ignored=True)

            opted_out_updated, activities_cancelled = _apply_opted_out_rule_existing_task(
                task_id=task_id,
                current_status=current_status,
            )

            uc_id = _ensure_use_case(data)
            _sync_use_case_exit_criteria(data, uc_id)

            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO row={row_number} opted-out rule applied task_id={task_id} ws={data.get('ws')} updated={opted_out_updated} activities_cancelled={activities_cancelled}",
                )

            return RowProcessResult(
                success=True,
                updated=opted_out_updated,
                cancelled=opted_out_updated,
                activities_cancelled=activities_cancelled,
            )

        if (
            data.get("task_eligible") == "N"
            and opt_in_label != "OPTED_IN"
            and current_eligible == "N"
            and current_status in CLOSED_TASK_STATUSES
        ):
            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO row={row_number} ignored not eligible already closed task_id={task_id} ws={data.get('ws')}",
                )
            return RowProcessResult(success=True, ignored=True)

        if (
            data.get("task_eligible") == "N"
            and opt_in_label != "OPTED_IN"
            and current_eligible == "Y"
            and current_status not in CLOSED_TASK_STATUSES
        ):
            activities_cancelled = _close_task_as_not_eligible(task_id=task_id)

            uc_id = _ensure_use_case(data)
            _sync_use_case_exit_criteria(data, uc_id)

            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO row={row_number} task closed as not eligible task_id={task_id} ws={data.get('ws')} activities_cancelled={activities_cancelled}",
                )

            return RowProcessResult(
                success=True,
                updated=True,
                cancelled=True,
                activities_cancelled=activities_cancelled,
            )

        task_reopened = False
        if (
            data.get("task_eligible") == "Y"
            and opt_in_label == "OPTED_IN"
            and current_status in CLOSED_TASK_STATUSES
        ):
            task_reopened = _reopen_task_if_needed(task_id=task_id)

        updated = _update_existing_task(
            task_id=task_id,
            existing_task=existing_task,
            payload=payload,
        )

        _sync_opt_in_status_history(
            task_id=task_id,
            opt_in_raw=data.get("opt_in_raw"),
            record_date=data.get("task_start_performed") or data.get("task_booking_date"),
            is_new_task=False,
        )

        activities_ok, activities_updated, activities_error, activities_error_column, activities_error_value, activity_count = _sync_task_activities(
            task_id=task_id,
            data=data,
            file_path=file_path,
            row_number=row_number,
        )

        if not activities_ok:
            _safe_log(
                file_path,
                row_number,
                activities_error or "Falha ao sincronizar activities",
                activities_error_column,
                activities_error_value,
            )
            return RowProcessResult(
                success=False,
                error_message=activities_error or "Falha ao sincronizar activities",
                error_column=activities_error_column,
                error_value=activities_error_value,
            )

        task_completed_updated = False
        if activity_count > 0:
            task_completed_updated = _recalculate_task_completed_if_needed(
                task_id=task_id,
                current_task_status=current_status,
                current_task_completed=existing_task.get("task_completed"),
            )

        uc_id = _ensure_use_case(data)
        _sync_use_case_exit_criteria(data, uc_id)

        if execution_log_path:
            _append_execution_log(
                execution_log_path,
                f"INFO row={row_number} task existing processed task_id={task_id} ws={data.get('ws')} updated={updated} reopened={task_reopened} activities_updated={activities_updated} task_completed_updated={task_completed_updated}",
            )

        return RowProcessResult(
            success=True,
            updated=(updated or task_reopened or activities_updated > 0 or task_completed_updated),
            ignored=not (updated or task_reopened or activities_updated > 0 or task_completed_updated),
            activities_updated=activities_updated,
        )

    except Exception as ex:
        message = f"Unexpected error: {ex}"
        _safe_log(file_path, row_number, message, None, None)

        if execution_log_path:
            _append_execution_log(
                execution_log_path,
                f"ERROR row={row_number} unexpected error={str(ex)[:1000]} traceback={traceback.format_exc()[:2000]}",
            )

        return RowProcessResult(
            success=False,
            error_message=message,
            error_column=None,
            error_value=None,
        )


# ======================================================================================
# PROCESSAMENTO DE CHUNK
# ======================================================================================

def _process_chunk(
    ws: Worksheet,
    idx_map: Dict[str, int],
    headers: List[str],
    start_row: int,
    end_row: int,
    file_path: str,
    failed_output_path: Path,
    execution_log_path: Optional[Path] = None,
    chunk_number: int = 0,
) -> Dict[str, int]:
    metrics = {
        "success": 0,
        "error": 0,
        "ignored": 0,
        "created": 0,
        "updated": 0,
        "cancelled": 0,
        "activities_updated": 0,
        "activities_cancelled": 0,
    }

    failed_rows_buffer: List[List[Any]] = []

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO chunk_start chunk={chunk_number} start_row={start_row} end_row={end_row}",
        )

    for r in range(end_row, start_row - 1, -1):
        row_cells = ws[r]
        raw_values = [cell.value for cell in row_cells]
        row_dict: Dict[str, Any] = {}

        for h, idx in idx_map.items():
            row_dict[h] = raw_values[idx] if idx < len(raw_values) else None

        result = _process_single_row(
            row_dict=row_dict,
            file_path=file_path,
            row_number=r,
            execution_log_path=execution_log_path,
        )

        if result.success:
            metrics["success"] += 1
            if result.ignored:
                metrics["ignored"] += 1
            if result.created:
                metrics["created"] += 1
            if result.updated:
                metrics["updated"] += 1
            if result.cancelled:
                metrics["cancelled"] += 1
            metrics["activities_updated"] += result.activities_updated
            metrics["activities_cancelled"] += result.activities_cancelled
        else:
            metrics["error"] += 1
            failed_rows_buffer.append(
                _row_dict_to_failed_excel_row(
                    headers=headers,
                    row_dict=row_dict,
                    original_row_number=r,
                    input_file_path=file_path,
                    result=result,
                )
            )

        ws.delete_rows(r, 1)

    if failed_rows_buffer:
        failed_rows_buffer.reverse()
        _append_failed_rows(
            failed_path=failed_output_path,
            rows_to_append=failed_rows_buffer,
            original_headers=headers,
            execution_log_path=execution_log_path,
        )

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO chunk_finish chunk={chunk_number} success={metrics['success']} error={metrics['error']} ignored={metrics['ignored']} created={metrics['created']} updated={metrics['updated']} cancelled={metrics['cancelled']} activities_updated={metrics['activities_updated']} activities_cancelled={metrics['activities_cancelled']}",
        )

    return metrics


# ======================================================================================
# FUNÇÃO PÚBLICA PRINCIPAL
# ======================================================================================

def run_import(
    file_path: str,
    user_id: Optional[str] = None,
    import_control_id: Optional[int] = None,
) -> Dict[str, Any]:
    _ensure_directories()

    input_path = _resolve_input_path(file_path)
    failed_output_path = _build_failed_output_path(input_path)
    execution_log_path = _build_execution_log_path(input_path)

    if import_control_id is None and repo_import_control is not None:
        try:
            import_control_id = repo_import_control.get_id_by_file(input_path.name)
        except Exception:
            import_control_id = None

    started_at = datetime.now()
    total_rows = 0
    chunk_number = 0

    summary_metrics = {
        "processed_success": 0,
        "failed_rows": 0,
        "ignored_rows": 0,
        "tasks_created": 0,
        "tasks_updated": 0,
        "tasks_cancelled": 0,
        "activities_updated": 0,
        "activities_cancelled": 0,
    }

    _append_execution_log(execution_log_path, "START import")
    _append_execution_log(execution_log_path, f"INFO import_source={IMPORT_SOURCE}")
    _append_execution_log(execution_log_path, f"INFO user_id={user_id}")
    _append_execution_log(execution_log_path, f"INFO input_file={str(input_path)}")
    _append_execution_log(execution_log_path, f"INFO failed_file={str(failed_output_path)}")

    try:
        wb, ws, idx_map, headers = _open_workbook_rw(str(input_path))
    except Exception as e:
        msg = f"Erro ao abrir arquivo Excel: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR failed opening workbook error={str(e)[:1000]}")
        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": 0,
                "processed_success": 0,
                "failed_rows": 0,
                "execution_log_path": str(execution_log_path),
            },
        }

    total_rows = max(ws.max_row - 1, 0)
    _append_execution_log(execution_log_path, f"INFO total_rows_initial={total_rows}")
    _append_execution_log(execution_log_path, f"INFO chunk_size={CHUNK_SIZE}")

    missing_cols = _validate_required_columns(headers)
    if missing_cols:
        msg = "Missing columns: " + ", ".join(missing_cols)
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR missing_columns={', '.join(missing_cols)}")
        wb.close()
        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": total_rows,
                "processed_success": 0,
                "failed_rows": total_rows,
                "execution_log_path": str(execution_log_path),
            },
        }

    try:
        while ws.max_row > 1:
            chunk_number += 1
            chunk_end = min(1 + CHUNK_SIZE, ws.max_row)

            chunk_metrics = _process_chunk(
                ws=ws,
                idx_map=idx_map,
                headers=headers,
                start_row=2,
                end_row=chunk_end,
                file_path=str(input_path),
                failed_output_path=failed_output_path,
                execution_log_path=execution_log_path,
                chunk_number=chunk_number,
            )

            summary_metrics["processed_success"] += chunk_metrics["success"]
            summary_metrics["failed_rows"] += chunk_metrics["error"]
            summary_metrics["ignored_rows"] += chunk_metrics["ignored"]
            summary_metrics["tasks_created"] += chunk_metrics["created"]
            summary_metrics["tasks_updated"] += chunk_metrics["updated"]
            summary_metrics["tasks_cancelled"] += chunk_metrics["cancelled"]
            summary_metrics["activities_updated"] += chunk_metrics["activities_updated"]
            summary_metrics["activities_cancelled"] += chunk_metrics["activities_cancelled"]

            wb.save(str(input_path))
            remaining_rows_current = max(ws.max_row - 1, 0)

            _append_execution_log(
                execution_log_path,
                f"INFO checkpoint_after_chunk chunk={chunk_number} success_total={summary_metrics['processed_success']} error_total={summary_metrics['failed_rows']} ignored_total={summary_metrics['ignored_rows']} remaining_rows_current={remaining_rows_current}",
            )

            if chunk_number % IMPORT_CONTROL_PROGRESS_INTERVAL_CHUNKS == 0:
                _report_progress(
                    import_control_id=import_control_id,
                    total_rows=total_rows,
                    processed_success=summary_metrics["processed_success"],
                    failed_rows=summary_metrics["failed_rows"],
                    ignored_rows=summary_metrics["ignored_rows"],
                    remaining_rows=remaining_rows_current,
                    execution_log_path=execution_log_path,
                )

    except Exception as e:
        msg = f"Erro durante o processamento da importação: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR processing failed error={str(e)[:1000]}")

        try:
            wb.save(str(input_path))
            _append_execution_log(execution_log_path, "INFO workbook saved after processing exception")
        except Exception as save_ex:
            _append_execution_log(
                execution_log_path,
                f"ERROR failed saving workbook after exception error={str(save_ex)[:1000]}",
            )

        wb.close()

        finished_at = datetime.now()
        duration_seconds = int((finished_at - started_at).total_seconds())

        _append_execution_log(
            execution_log_path,
            f"FINISH status=FAILED success={summary_metrics['processed_success']} error={summary_metrics['failed_rows']} duration_seconds={duration_seconds}",
        )

        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": total_rows,
                "processed_success": summary_metrics["processed_success"],
                "failed_rows": summary_metrics["failed_rows"],
                "ignored_rows": summary_metrics["ignored_rows"],
                "tasks_created": summary_metrics["tasks_created"],
                "tasks_updated": summary_metrics["tasks_updated"],
                "tasks_cancelled": summary_metrics["tasks_cancelled"],
                "activities_updated": summary_metrics["activities_updated"],
                "activities_cancelled": summary_metrics["activities_cancelled"],
                "failed_file_path": str(failed_output_path),
                "execution_log_path": str(execution_log_path),
            },
        }

    try:
        wb.save(str(input_path))
        _append_execution_log(execution_log_path, "INFO workbook saved successfully at end")
    except Exception as e:
        msg = f"Erro ao salvar arquivo Excel após processamento: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR final workbook save failed error={str(e)[:1000]}")
        wb.close()

        finished_at = datetime.now()
        duration_seconds = int((finished_at - started_at).total_seconds())

        _append_execution_log(
            execution_log_path,
            f"FINISH status=FAILED success={summary_metrics['processed_success']} error={summary_metrics['failed_rows']} duration_seconds={duration_seconds}",
        )

        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": total_rows,
                "processed_success": summary_metrics["processed_success"],
                "failed_rows": summary_metrics["failed_rows"],
                "ignored_rows": summary_metrics["ignored_rows"],
                "tasks_created": summary_metrics["tasks_created"],
                "tasks_updated": summary_metrics["tasks_updated"],
                "tasks_cancelled": summary_metrics["tasks_cancelled"],
                "activities_updated": summary_metrics["activities_updated"],
                "activities_cancelled": summary_metrics["activities_cancelled"],
                "failed_file_path": str(failed_output_path),
                "execution_log_path": str(execution_log_path),
            },
        }

    wb.close()

    remaining_rows_in_input = 0
    try:
        wb_check = load_workbook(str(input_path), read_only=True, data_only=True)
        ws_check = wb_check.active
        remaining_rows_in_input = max(ws_check.max_row - 1, 0)
        wb_check.close()
    except Exception as e:
        remaining_rows_in_input = -1
        _append_execution_log(
            execution_log_path,
            f"WARN failed checking remaining rows error={str(e)[:1000]}",
        )

    finished_at = datetime.now()
    duration_seconds = int((finished_at - started_at).total_seconds())

    _append_execution_log(execution_log_path, f"INFO total_chunks={chunk_number}")
    _append_execution_log(execution_log_path, f"INFO total_rows_initial={total_rows}")
    _append_execution_log(execution_log_path, f"INFO processed_success={summary_metrics['processed_success']}")
    _append_execution_log(execution_log_path, f"INFO failed_rows={summary_metrics['failed_rows']}")
    _append_execution_log(execution_log_path, f"INFO ignored_rows={summary_metrics['ignored_rows']}")
    _append_execution_log(execution_log_path, f"INFO tasks_created={summary_metrics['tasks_created']}")
    _append_execution_log(execution_log_path, f"INFO tasks_updated={summary_metrics['tasks_updated']}")
    _append_execution_log(execution_log_path, f"INFO tasks_cancelled={summary_metrics['tasks_cancelled']}")
    _append_execution_log(execution_log_path, f"INFO activities_updated={summary_metrics['activities_updated']}")
    _append_execution_log(execution_log_path, f"INFO activities_cancelled={summary_metrics['activities_cancelled']}")
    _append_execution_log(execution_log_path, f"INFO remaining_rows_in_input={remaining_rows_in_input}")
    _append_execution_log(execution_log_path, f"INFO failed_file={str(failed_output_path)}")
    _append_execution_log(execution_log_path, f"INFO duration_seconds={duration_seconds}")

    if summary_metrics["failed_rows"] > 0:
        status = "FAILED"
        msg = (
            f"{IMPORT_SOURCE} import concluído com falhas. "
            f"arquivo={str(input_path)}, total={total_rows}, sucesso={summary_metrics['processed_success']}, "
            f"erros={summary_metrics['failed_rows']}, ignoradas={summary_metrics['ignored_rows']}, "
            f"criadas={summary_metrics['tasks_created']}, atualizadas={summary_metrics['tasks_updated']}, "
            f"canceladas={summary_metrics['tasks_cancelled']}, activities_atualizadas={summary_metrics['activities_updated']}, "
            f"activities_cancelled={summary_metrics['activities_cancelled']}, linhas_restantes_input={remaining_rows_in_input}, "
            f"arquivo_falhas={str(failed_output_path)}, log_execucao={str(execution_log_path)}."
        )
    else:
        status = "FINISHED"
        msg = (
            f"{IMPORT_SOURCE} import concluído com sucesso. "
            f"arquivo={str(input_path)}, total={total_rows}, sucesso={summary_metrics['processed_success']}, "
            f"erros={summary_metrics['failed_rows']}, ignoradas={summary_metrics['ignored_rows']}, "
            f"criadas={summary_metrics['tasks_created']}, atualizadas={summary_metrics['tasks_updated']}, "
            f"canceladas={summary_metrics['tasks_cancelled']}, activities_atualizadas={summary_metrics['activities_updated']}, "
            f"activities_cancelled={summary_metrics['activities_cancelled']}, linhas_restantes_input={remaining_rows_in_input}, "
            f"arquivo_falhas={str(failed_output_path)}, log_execucao={str(execution_log_path)}."
        )

    _append_execution_log(
        execution_log_path,
        f"FINISH status={status} success={summary_metrics['processed_success']} error={summary_metrics['failed_rows']} ignored={summary_metrics['ignored_rows']} remaining_rows_in_input={remaining_rows_in_input} duration_seconds={duration_seconds}",
    )

    return {
        "status": status,
        "message": msg,
        "summary": {
            "file_path": str(input_path),
            "total_rows": total_rows,
            "processed_success": summary_metrics["processed_success"],
            "failed_rows": summary_metrics["failed_rows"],
            "ignored_rows": summary_metrics["ignored_rows"],
            "tasks_created": summary_metrics["tasks_created"],
            "tasks_updated": summary_metrics["tasks_updated"],
            "tasks_cancelled": summary_metrics["tasks_cancelled"],
            "activities_updated": summary_metrics["activities_updated"],
            "activities_cancelled": summary_metrics["activities_cancelled"],
            "failed_file_path": str(failed_output_path),
            "remaining_rows_in_input": remaining_rows_in_input,
            "execution_log_path": str(execution_log_path),
            "duration_seconds": duration_seconds,
            "total_chunks": chunk_number,
        },
    }
