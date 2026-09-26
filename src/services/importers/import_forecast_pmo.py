"""
import_forecast_pmo.py

Importador Python do Forecast PMO.

Objetivo
--------
Converter o processo legado em VBA para o padrão atual de importadores Python
do Bridge Adoption, com foco em:

- leitura direta de arquivo XLSX;
- processamento em chunks;
- remoção das linhas já lidas do arquivo original;
- geração de arquivo XLSX de falhas;
- geração de log texto de execução;
- gravação funcional em tbImportLog quando disponível;
- baixo impacto de memória/CPU no servidor.

Regras de negócio consolidadas
------------------------------
1. O identificador funcional do projeto no PMO é a coluna OV.
2. A OV pode ser composta por várias OVs separadas por "_".
3. A OV pode conter "#" como marcação de Oracle Fusion; em tbProjectOV
   devem ser guardados os itens individualizados e sem "#".
4. O importador usa ProjectRepository.insert() como núcleo para insert/update
   de tbProject + sincronização de tbProjectOV.
5. Projetos sem OV numérica também são válidos (POC, PSR, PSR-..., etc.).
6. Pode existir mais de uma linha do mesmo projeto por causa de TP.
   Neste caso:
   - a primeira linha válida é a linha principal do projeto;
   - as demais linhas servem para sincronizar TP e equipe;
   - o project_id deve ser o mesmo;
   - não deve haver duplicidade em tbProjectTeam.
7. Se o cliente não existir em tbCompanyListName:
   - a linha falha;
   - todas as linhas da mesma OV também devem ir para o XLSX de erros.
8. Se uma pessoa não existir em tbPersonListName:
   - o projeto continua sendo importado;
   - a pessoa é ignorada na equipe;
   - a(s) linha(s) da mesma OV vai(vão) para o XLSX de erros
     indicando a(s) pessoa(s) faltante(s).
9. Projetos já fechados/cancelados na base não são reabertos/alterados,
   mantendo a regra do VBA.
10. Baseline/progress comentado no VBA fica fora do escopo.
11. Cabeçalho do Excel não é fixo:
   - deve ser localizado dinamicamente;
   - os dados começam logo abaixo do cabeçalho válido.

Observações de alocação
-----------------------
- allocation_start: Internalization Date; se ausente/ inválida, usa fallback
  equivalente ao legado: dia anterior ao início do FY corrente (31/03 do FY).
- allocation_end: Project Closure Date quando houver.
- para projeto Closed/Canceled com data de fechamento, recursos ativos do
  projeto são encerrados em tbProjectTeam.
"""

from __future__ import annotations

import logging
import re
import traceback
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.infrastructure.database.repositories.company_list_name_repository import (
    CompanyListNameRepository,
)
from src.infrastructure.database.repositories.person_list_name_repository import (
    PersonListNameRepository,
)
from src.infrastructure.database.repositories.project_repository import ProjectRepository
from src.infrastructure.database.repositories.project_team_repository import (
    ProjectTeamRepository,
)
from src.infrastructure.database.repositories.project_type_repository import (
    ProjectTypeRepository,
)
from src.infrastructure.database.repositories.squad_repository import SquadRepository

try:
    from src.infrastructure.database.repositories.import_log_repository import ImportLogRepository
except Exception:
    ImportLogRepository = None

try:
    from src.infrastructure.database.repositories.import_control_repository import (
        ImportControlRepository,
        ImportStatus,
    )
except Exception:
    ImportControlRepository = None
    ImportStatus = None


BASE_STORAGE_PATH = Path("/home/bridgeadoption/storage")
BASE_INPUT_PATH = BASE_STORAGE_PATH / "input"
BASE_OUTPUT_PATH = BASE_STORAGE_PATH / "output"
BASE_LOGS_PATH = BASE_STORAGE_PATH / "logs"

IMPORT_SOURCE = "ImportForecastPMO"

ERROR_EXTRA_COLUMNS = [
    "import_error_message",
    "import_error_column",
    "import_error_value",
    "import_original_row",
    "import_processed_at",
    "import_source",
    "import_original_file",
]

CHUNK_SIZE = 200
IMPORT_CONTROL_PROGRESS_INTERVAL_CHUNKS = 1

PROJECT_STATUS_MAP = {
    "em andamento": "In progress",
    "andamento": "In progress",
    "não iniciado": "Not started",
    "nao iniciado": "Not started",
    "iniciar": "Not started",
    "on-hold": "On Hold",
    "encerrado": "Closed",
    "concluído": "Closed",
    "concluido": "Closed",
    "cancelado": "Canceled",
    "modelo de negócio": "Business Model",
    "modelo de negocio": "Business Model",
    "avaliação": "Under Analysis",
    "avaliacao": "Under Analysis",
    "em avaliação": "Under Analysis",
    "em avaliacao": "Under Analysis",
}

CLOSED_PROJECT_STATUSES = {"Closed", "Canceled"}

ROLE_CONFIG = {
    "PM": {
        "column": "PM",
        "default_department_id": 11,
        "default_level_id": 2,
        "technical_lead": 0,
    },
    "AM": {
        "column": "AM",
        "default_department_id": 35,
        "default_level_id": 9,
        "technical_lead": 0,
    },
    "PRE_SALES": {
        "column": "Pre Sales",
        "default_department_id": 23,
        "default_level_id": 10,
        "technical_lead": 0,
    },
    "TECHNICAL_LEAD": {
        "column": "Technical Lead",
        "default_department_id": 0,
        "default_level_id": 0,
        "technical_lead": 1,
    },
}

REQUIRED_HEADERS = [
    "OV",
    "TP",
    "Client",
    "Project Name",
    "PM",
    "AM",
    "Internalization Date",
    "Project Start Date",
    "Project Closure Date",
    "Status",
    "Pre Sales",
    "Technical Lead",
]

HEADER_ANCHORS = ("OV", "Project Name")

logger = logging.getLogger(__name__)

repo_company = CompanyListNameRepository()
repo_person_list = PersonListNameRepository()
repo_project = ProjectRepository()
repo_project_team = ProjectTeamRepository()
repo_project_type = ProjectTypeRepository()
repo_squad = SquadRepository()
repo_log = ImportLogRepository() if ImportLogRepository else None
repo_import_control = ImportControlRepository() if ImportControlRepository else None


@dataclass
class RowProcessResult:
    success: bool
    error_message: Optional[str] = None
    error_column: Optional[str] = None
    error_value: Optional[Any] = None
    warnings: List[str] = field(default_factory=list)
    group_fail_ov: bool = False
    project_id: int = 0
    created: bool = False
    updated: bool = False
    ignored: bool = False


def _ensure_directories() -> None:
    BASE_INPUT_PATH.mkdir(parents=True, exist_ok=True)
    BASE_OUTPUT_PATH.mkdir(parents=True, exist_ok=True)
    BASE_LOGS_PATH.mkdir(parents=True, exist_ok=True)


def _resolve_input_path(file_path: str) -> Path:
    path = Path(file_path)
    if path.is_absolute():
        return path
    return BASE_INPUT_PATH / path.name


def _build_failed_output_path(input_path: Path) -> Path:
    return BASE_OUTPUT_PATH / f"{input_path.stem}_failed_rows.xlsx"


def _build_execution_log_path(input_path: Path) -> Path:
    return BASE_LOGS_PATH / f"{input_path.stem}.log"


def _append_execution_log(log_path: Path, message: str) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{ts}] {message}\n")


def _build_row_log_message(
    row_number: int,
    ov: Optional[str],
    result: RowProcessResult,
) -> str:
    ov_value = ov or "<empty>"
    base_parts = [f"row={row_number}", f"ov={ov_value}"]

    if result.success:
        if result.ignored:
            base_parts.append("result=ignored")
            if result.project_id:
                base_parts.append(f"project_id={result.project_id}")
            return "INFO " + " ".join(base_parts) + ' reason="project already closed"'

        if result.warnings:
            base_parts.append("result=warning")
        else:
            base_parts.append("result=success")

        if result.project_id:
            base_parts.append(f"project_id={result.project_id}")

        base_parts.append(f"created={1 if result.created else 0}")
        base_parts.append(f"updated={1 if result.updated else 0}")
        base_parts.append(f"warnings={len(result.warnings)}")

        if result.warnings:
            warning_message = " ; ".join(result.warnings)[:2000]
            return "WARN " + " ".join(base_parts) + f' message="{warning_message}"'

        return "INFO " + " ".join(base_parts)

    base_parts.append("result=error")
    if result.error_column:
        base_parts.append(f"column={result.error_column}")
    if result.error_value is not None:
        safe_value = str(result.error_value).replace('"', "'")[:500]
        base_parts.append(f'value="{safe_value}"')

    safe_message = str(result.error_message or "Unknown error").replace('"', "'")[:2000]
    return "ERROR " + " ".join(base_parts) + f' message="{safe_message}"'


def _safe_log(
    file_path: str,
    row_number: int,
    message: str,
    column_name: Optional[str] = None,
    value: Optional[Any] = None,
) -> None:
    if repo_log is None:
        return

    try:
        repo_log.create(
            IMPORT_SOURCE,
            file_path,
            row_number,
            message,
            column_name,
            value,
        )
    except Exception as ex:
        logger.exception(
            "Falha ao gravar tbImportLog. file=%s row=%s column=%s error=%s",
            file_path,
            row_number,
            column_name,
            ex,
        )


def _report_progress(
    import_control_id: Optional[int],
    total_rows: int,
    processed_success: int,
    failed_rows: int,
    ignored_rows: int,
    remaining_rows: int,
    execution_log_path: Optional[Path] = None,
) -> None:
    if repo_import_control is None or import_control_id is None:
        return

    lidas = total_rows - remaining_rows
    msg = (
        f"Arquivo processado. Total origem={total_rows} | "
        f"lidas={lidas} | sucesso={processed_success} | "
        f"erros={failed_rows} | ignoradas={ignored_rows} | restantes={remaining_rows}"
    )

    try:
        repo_import_control.update_status(
            importctrl_id=import_control_id,
            status=ImportStatus.RUNNING,
            message=msg,
        )
    except Exception as e:
        if execution_log_path:
            _append_execution_log(
                execution_log_path,
                f"WARN failed updating import control progress error={str(e)[:500]}",
            )


def _normalize_str(value: Any, max_length: int = 1000) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text[:max_length]


def _normalize_header(value: Any) -> str:
    text = _normalize_str(value, max_length=200) or ""
    return re.sub(r"\s+", " ", text).strip().lower()


def _normalize_status(value: Any) -> str:
    raw = _normalize_str(value, max_length=200)
    if not raw:
        return "Unidentified"
    key = raw.lower()
    return PROJECT_STATUS_MAP.get(key, "Unidentified")


def _to_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()[:50]
    if not text:
        return None

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d/%m/%y",
        "%m/%d/%y",
        "%d.%m.%Y",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def _current_fy_internalization_fallback() -> date:
    today = date.today()
    fy_start_year = today.year if today.month >= 4 else today.year - 1
    return date(fy_start_year, 3, 31)


def _split_people(raw_value: Any) -> List[str]:
    text = _normalize_str(raw_value, max_length=4000)
    if not text:
        return []

    normalized = text
    normalized = normalized.replace("\r\n", ";")
    normalized = normalized.replace("\n", ";")
    normalized = normalized.replace("\r", ";")

    for token in ["/", "|", ",", "+", "&"]:
        normalized = normalized.replace(token, ";")

    normalized = re.sub(r"\s+e\s+", ";", normalized, flags=re.IGNORECASE)

    parts = []
    seen = set()
    for item in normalized.split(";"):
        person = str(item).strip()
        if not person:
            continue
        if person.upper() in {"N/A", "NA", "TBD", "0", "-"}:
            continue
        if len(person) <= 1:
            continue
        if person not in seen:
            seen.add(person)
            parts.append(person)

    return parts


def _find_header_row(ws: Worksheet) -> Tuple[int, Dict[str, int], List[str]]:
    max_scan_rows = min(ws.max_row, 30)
    best_row = 0
    best_headers: List[str] = []
    best_idx_map: Dict[str, int] = {}
    best_required_count = -1

    normalized_required = {_normalize_header(x): x for x in REQUIRED_HEADERS}

    for r in range(1, max_scan_rows + 1):
        row_values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        normalized_cells = [_normalize_header(v) for v in row_values]

        if not any(anchor.lower() in normalized_cells for anchor in [a.lower() for a in HEADER_ANCHORS]):
            continue

        idx_map: Dict[str, int] = {}
        matched_required = 0
        original_headers: List[str] = []

        for idx, cell_value in enumerate(row_values, start=1):
            header_text = _normalize_str(cell_value, max_length=200) or ""
            original_headers.append(header_text)
            normalized = _normalize_header(cell_value)

            if normalized in normalized_required and normalized_required[normalized] not in idx_map:
                idx_map[normalized_required[normalized]] = idx
                matched_required += 1

        if matched_required > best_required_count:
            best_required_count = matched_required
            best_row = r
            best_headers = original_headers
            best_idx_map = idx_map

    if best_row == 0:
        raise ValueError("Não foi possível localizar um cabeçalho válido no arquivo XLSX.")

    missing = [h for h in REQUIRED_HEADERS if h not in best_idx_map]
    if missing:
        raise ValueError(f"Cabeçalho encontrado, porém faltam colunas obrigatórias: {', '.join(missing)}")

    return best_row, best_idx_map, best_headers


def _open_workbook_rw(file_path: str) -> Tuple[Any, Worksheet, int, Dict[str, int], List[str]]:
    wb = load_workbook(filename=file_path, read_only=False, data_only=True)
    ws = wb.active
    header_row, idx_map, headers = _find_header_row(ws)
    return wb, ws, header_row, idx_map, headers


def _ensure_failed_workbook(
    failed_path: Path,
    original_headers: List[str],
    execution_log_path: Optional[Path] = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "failed_rows"
    ws.append(original_headers + ERROR_EXTRA_COLUMNS)
    wb.save(str(failed_path))
    wb.close()

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO failed file created (overwritten if existed): {str(failed_path)}",
        )


def _append_failed_rows(
    failed_path: Path,
    rows_to_append: List[List[Any]],
    original_headers: List[str],
    execution_log_path: Optional[Path] = None,
) -> None:
    if not rows_to_append:
        return

    if not failed_path.exists():
        _ensure_failed_workbook(failed_path, original_headers, execution_log_path)

    wb = load_workbook(str(failed_path))
    ws = wb.active
    for row in rows_to_append:
        ws.append(row)
    wb.save(str(failed_path))
    wb.close()

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO appended failed rows count={len(rows_to_append)} to file={str(failed_path)}",
        )


def _row_dict_to_failed_excel_row(
    headers: List[str],
    row_dict: Dict[str, Any],
    original_row_number: int,
    input_file_path: str,
    result: RowProcessResult,
) -> List[Any]:
    processed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    original_values = [row_dict.get(h) for h in headers]
    if result.warnings:
        warning_suffix = " | WARNINGS: " + " ; ".join(result.warnings)
        error_message = (result.error_message or "Linha processada com advertências") + warning_suffix
    else:
        error_message = result.error_message

    return original_values + [
        error_message,
        result.error_column,
        str(result.error_value)[:2000] if result.error_value is not None else None,
        original_row_number,
        processed_at,
        IMPORT_SOURCE,
        input_file_path,
    ]


def _get_row_dict_from_ws(
    ws: Worksheet,
    row_number: int,
    idx_map: Dict[str, int],
    headers: List[str],
) -> Dict[str, Any]:
    row_dict: Dict[str, Any] = {h: None for h in headers}
    for header_name, column_index in idx_map.items():
        row_dict[header_name] = ws.cell(row_number, column_index).value
    return row_dict


def _normalize_ov_key(value: Any) -> Optional[str]:
    text = _normalize_str(value, max_length=500)
    if not text:
        return None
    return re.sub(r"\s+", "", text)


def _resolve_company_id(company_name: Optional[str]) -> int:
    if not company_name:
        return 0
    try:
        result = repo_company.get_company_id_by_name(company_name)
        return int(result or 0)
    except Exception:
        return 0


def _resolve_person_id(name: str) -> int:
    try:
        result = repo_person_list.get_person_id_by_name(name)
        return int(result or 0)
    except Exception:
        return 0


def _resolve_latest_squad_for_person(person_id: int) -> Optional[Dict[str, Any]]:
    if not person_id:
        return None

    try:
        ids = repo_squad.find_ids_by({"squad_person_id": person_id})
    except Exception:
        ids = []

    if not ids:
        return None

    rows = repo_squad.get_squad(squad_id=ids, as_df=False)
    if not rows:
        return None

    def squad_sort_key(item: Dict[str, Any]) -> Tuple[date, int]:
        upgrade = _to_date(item.get("squad_upgrade")) or date(1900, 1, 1)
        squad_id = int(item.get("squad_id", 0) or 0)
        return (upgrade, squad_id)

    rows_sorted = sorted(rows, key=squad_sort_key, reverse=True)
    return rows_sorted[0] if rows_sorted else None


def _resolve_role_department_level(person_id: int, role_key: str) -> Tuple[int, int]:
    config = ROLE_CONFIG[role_key]
    latest_squad = _resolve_latest_squad_for_person(person_id)

    if latest_squad:
        dept_id = int(latest_squad.get("squad_department_id") or 0)
        level_id = int(latest_squad.get("squad_level_id") or 0)
        if dept_id > 0 and level_id > 0:
            return dept_id, level_id

    return config["default_department_id"], config["default_level_id"]


def _get_existing_project_status(project_id: int) -> Optional[str]:
    rows = repo_project.get_project(as_df=False)
    for row in rows:
        if int(row.get("project_id", 0) or 0) == project_id:
            return row.get("project_status")
    return None


def _build_project_payload_from_row(row_dict: Dict[str, Any], company_id: int) -> Dict[str, Any]:
    ov = _normalize_ov_key(row_dict.get("OV"))
    project_name = _normalize_str(row_dict.get("Project Name"), max_length=1000)
    client_name = _normalize_str(row_dict.get("Client"), max_length=1000)
    normalized_status = _normalize_status(row_dict.get("Status"))
    internalization_date = _to_date(row_dict.get("Internalization Date")) or _current_fy_internalization_fallback()
    start_date = _to_date(row_dict.get("Project Start Date"))
    end_date = _to_date(row_dict.get("Project Closure Date"))

    payload: Dict[str, Any] = {
        "project_ov": ov,
        "project_customer_id": company_id,
        "project_customer_name": client_name,
        "project_name": project_name,
        "project_owner": "PMO",
        "project_status": normalized_status,
        "project_internalization_date": internalization_date,
    }

    if start_date is not None:
        payload["project_start_date"] = start_date

    if end_date is not None:
        payload["project_end_date"] = end_date

    return payload


def _sync_project_team_from_row(
    project_id: int,
    row_dict: Dict[str, Any],
    warnings: List[str],
    execution_log_path: Optional[Path] = None,
) -> None:
    allocation_start = _to_date(row_dict.get("Internalization Date")) or _current_fy_internalization_fallback()
    allocation_end = _to_date(row_dict.get("Project Closure Date"))

    for role_key, role_cfg in ROLE_CONFIG.items():
        names = _split_people(row_dict.get(role_cfg["column"]))
        for person_name in names:
            person_id = _resolve_person_id(person_name)
            if not person_id:
                warnings.append(f"{role_cfg['column']}: pessoa não encontrada '{person_name}'")
                continue

            department_id, level_id = _resolve_role_department_level(person_id, role_key)
            technical_lead = int(role_cfg["technical_lead"])

            allocated_id = repo_project_team.allocate_person(
                project_id=project_id,
                person_id=person_id,
                department_id=department_id,
                level_id=level_id,
                allocation_start=allocation_start,
                allocation_end=allocation_end,
                technical_lead=technical_lead,
            )

            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO project_team role={role_key} project_id={project_id} person_id={person_id} allocated_id={allocated_id} technical_lead={technical_lead}",
                )


def _close_project_team_if_needed(project_id: int, row_dict: Dict[str, Any]) -> int:
    status = _normalize_status(row_dict.get("Status"))
    closure_date = _to_date(row_dict.get("Project Closure Date"))
    if status not in CLOSED_PROJECT_STATUSES or closure_date is None:
        return 0
    return repo_project_team.close_active_rows_by_project(project_id, closure_date)


def _process_single_row(
    row_dict: Dict[str, Any],
    file_path: str,
    row_number: int,
    ov_group_rows_count: int,
    first_row_for_ov: bool,
    processed_projects_by_ov: Dict[str, int],
    execution_log_path: Optional[Path] = None,
) -> RowProcessResult:
    try:
        ov = _normalize_ov_key(row_dict.get("OV"))
        if not ov:
            message = "Campo obrigatório ausente ou inválido: OV"
            _safe_log(file_path, row_number, message, "OV", row_dict.get("OV"))
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="OV",
                error_value=row_dict.get("OV"),
            )

        client_name = _normalize_str(row_dict.get("Client"), max_length=1000)
        if not client_name:
            message = "Campo obrigatório ausente ou inválido: Client"
            _safe_log(file_path, row_number, message, "Client", row_dict.get("Client"))
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="Client",
                error_value=row_dict.get("Client"),
                group_fail_ov=True,
            )

        company_id = _resolve_company_id(client_name)
        if not company_id:
            message = "Cliente não encontrado em tbCompanyListName"
            _safe_log(file_path, row_number, message, "Client", client_name)
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="Client",
                error_value=client_name,
                group_fail_ov=True,
            )

        project_name = _normalize_str(row_dict.get("Project Name"), max_length=1000)
        if not project_name:
            message = "Campo obrigatório ausente ou inválido: Project Name"
            _safe_log(file_path, row_number, message, "Project Name", row_dict.get("Project Name"))
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="Project Name",
                error_value=row_dict.get("Project Name"),
                group_fail_ov=True,
            )

        project_id = int(processed_projects_by_ov.get(ov) or 0)
        created = False
        updated = False
        ignored = False

        if first_row_for_ov or project_id == 0:
            payload = _build_project_payload_from_row(row_dict, company_id)
            project_id = repo_project.insert(payload)
            if not project_id:
                message = "Falha ao inserir/atualizar projeto em tbProject"
                _safe_log(file_path, row_number, message, "OV", ov)
                return RowProcessResult(
                    success=False,
                    error_message=message,
                    error_column="OV",
                    error_value=ov,
                    group_fail_ov=True,
                )

            processed_projects_by_ov[ov] = project_id

            existing_status = _get_existing_project_status(project_id)
            if existing_status in CLOSED_PROJECT_STATUSES:
                ignored = True
            else:
                created = ov_group_rows_count == 1
                updated = not created

        if not project_id:
            message = "project_id inválido após processamento"
            _safe_log(file_path, row_number, message, "OV", ov)
            return RowProcessResult(
                success=False,
                error_message=message,
                error_column="OV",
                error_value=ov,
                group_fail_ov=True,
            )

        project_status = _get_existing_project_status(project_id)
        if project_status in CLOSED_PROJECT_STATUSES:
            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO row={row_number} ignored because project already closed project_id={project_id} ov={ov} status={project_status}",
                )
            return RowProcessResult(
                success=True,
                project_id=project_id,
                ignored=True,
            )

        project_type = _normalize_str(row_dict.get("TP"), max_length=200)
        if project_type:
            type_id = repo_project_type.ensure_type(project_id, project_type)
            if execution_log_path:
                _append_execution_log(
                    execution_log_path,
                    f"INFO project_type project_id={project_id} type={project_type} projecttp_id={type_id}",
                )

        warnings: List[str] = []
        _sync_project_team_from_row(
            project_id=project_id,
            row_dict=row_dict,
            warnings=warnings,
            execution_log_path=execution_log_path,
        )

        closed_rows = _close_project_team_if_needed(project_id, row_dict)
        if execution_log_path and closed_rows > 0:
            _append_execution_log(
                execution_log_path,
                f"INFO project_team_closed project_id={project_id} closed_rows={closed_rows}",
            )

        if warnings:
            for warning in warnings:
                _safe_log(file_path, row_number, warning, None, ov)

        return RowProcessResult(
            success=True,
            warnings=warnings,
            project_id=project_id,
            created=created,
            updated=updated,
            ignored=ignored,
        )

    except Exception as ex:
        message = f"Unexpected error: {ex}"
        _safe_log(file_path, row_number, message, None, None)

        if execution_log_path:
            _append_execution_log(
                execution_log_path,
                f"ERROR row={row_number} unexpected error={str(ex)[:1000]} traceback={traceback.format_exc()[:2000]}",
            )

        return RowProcessResult(
            success=False,
            error_message=message,
        )


def _process_chunk(
    ws: Worksheet,
    idx_map: Dict[str, int],
    headers: List[str],
    start_row: int,
    end_row: int,
    file_path: str,
    failed_output_path: Path,
    processed_projects_by_ov: Dict[str, int],
    execution_log_path: Optional[Path] = None,
    chunk_number: int = 0,
) -> Dict[str, int]:
    metrics = {
        "success": 0,
        "error": 0,
        "ignored": 0,
        "created": 0,
        "updated": 0,
        "warning_rows": 0,
    }

    source_rows: List[Tuple[int, Dict[str, Any], str]] = []
    for r in range(start_row, end_row + 1):
        row_dict = _get_row_dict_from_ws(ws, r, idx_map, headers)
        ov = _normalize_ov_key(row_dict.get("OV")) or f"__ROW_{r}"
        source_rows.append((r, row_dict, ov))

    ov_groups: Dict[str, List[Tuple[int, Dict[str, Any]]]] = {}
    for row_num, row_dict, ov in source_rows:
        ov_groups.setdefault(ov, []).append((row_num, row_dict))

    failed_rows_buffer: List[List[Any]] = []

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO chunk_start chunk={chunk_number} start_row={start_row} end_row={end_row}",
        )

    rows_to_delete = 0

    for ov, rows in ov_groups.items():
        group_failed = False
        group_failure_result: Optional[RowProcessResult] = None

        sorted_rows = sorted(rows, key=lambda x: x[0])
        total_group_rows = len(sorted_rows)

        for index_in_group, (original_row_number, row_dict) in enumerate(sorted_rows):
            result = _process_single_row(
                row_dict=row_dict,
                file_path=file_path,
                row_number=original_row_number,
                ov_group_rows_count=total_group_rows,
                first_row_for_ov=index_in_group == 0,
                processed_projects_by_ov=processed_projects_by_ov,
                execution_log_path=execution_log_path,
            )

            if not result.success and result.group_fail_ov:
                group_failed = True
                group_failure_result = result
                if execution_log_path:
                    _append_execution_log(
                        execution_log_path,
                        _build_row_log_message(
                            row_number=original_row_number,
                            ov=ov,
                            result=result,
                        ),
                    )
                break

            if result.success:
                metrics["success"] += 1
                if result.ignored:
                    metrics["ignored"] += 1
                if result.created:
                    metrics["created"] += 1
                if result.updated:
                    metrics["updated"] += 1
                if result.warnings:
                    metrics["warning_rows"] += 1
                    failed_rows_buffer.append(
                        _row_dict_to_failed_excel_row(
                            headers=headers,
                            row_dict=row_dict,
                            original_row_number=original_row_number,
                            input_file_path=file_path,
                            result=result,
                        )
                    )

                if execution_log_path:
                    _append_execution_log(
                        execution_log_path,
                        _build_row_log_message(
                            row_number=original_row_number,
                            ov=ov,
                            result=result,
                        ),
                    )
            else:
                metrics["error"] += 1
                failed_rows_buffer.append(
                    _row_dict_to_failed_excel_row(
                        headers=headers,
                        row_dict=row_dict,
                        original_row_number=original_row_number,
                        input_file_path=file_path,
                        result=result,
                    )
                )

                if execution_log_path:
                    _append_execution_log(
                        execution_log_path,
                        _build_row_log_message(
                            row_number=original_row_number,
                            ov=ov,
                            result=result,
                        ),
                    )

        if group_failed and group_failure_result is not None:
            metrics["error"] += total_group_rows
            for original_row_number, row_dict in sorted_rows:
                failed_rows_buffer.append(
                    _row_dict_to_failed_excel_row(
                        headers=headers,
                        row_dict=row_dict,
                        original_row_number=original_row_number,
                        input_file_path=file_path,
                        result=group_failure_result,
                    )
                )

                if execution_log_path and original_row_number != sorted_rows[0][0]:
                    _append_execution_log(
                        execution_log_path,
                        _build_row_log_message(
                            row_number=original_row_number,
                            ov=ov,
                            result=group_failure_result,
                        ),
                    )

        rows_to_delete += total_group_rows

    for _ in range(rows_to_delete):
        ws.delete_rows(start_row, 1)

    if failed_rows_buffer:
        _append_failed_rows(
            failed_path=failed_output_path,
            rows_to_append=failed_rows_buffer,
            original_headers=headers,
            execution_log_path=execution_log_path,
        )

    if execution_log_path:
        _append_execution_log(
            execution_log_path,
            f"INFO chunk_finish chunk={chunk_number} success={metrics['success']} error={metrics['error']} ignored={metrics['ignored']} created={metrics['created']} updated={metrics['updated']} warning_rows={metrics['warning_rows']}",
        )

    return metrics


def run_import(
    file_path: str,
    user_id: Optional[str] = None,
    import_control_id: Optional[int] = None,
) -> Dict[str, Any]:
    _ensure_directories()

    input_path = _resolve_input_path(file_path)
    failed_output_path = _build_failed_output_path(input_path)
    execution_log_path = _build_execution_log_path(input_path)

    if import_control_id is None and repo_import_control is not None:
        try:
            import_control_id = repo_import_control.get_id_by_file(input_path.name)
        except Exception:
            import_control_id = None

    started_at = datetime.now()
    processed_projects_by_ov: Dict[str, int] = {}
    total_rows = 0
    chunk_number = 0

    summary_metrics = {
        "processed_success": 0,
        "failed_rows": 0,
        "ignored_rows": 0,
        "projects_created": 0,
        "projects_updated": 0,
        "warning_rows": 0,
    }

    _append_execution_log(execution_log_path, "START import")
    _append_execution_log(execution_log_path, f"INFO import_source={IMPORT_SOURCE}")
    _append_execution_log(execution_log_path, f"INFO user_id={user_id}")
    _append_execution_log(execution_log_path, f"INFO input_file={str(input_path)}")
    _append_execution_log(execution_log_path, f"INFO failed_file={str(failed_output_path)}")
    _append_execution_log(execution_log_path, f"INFO chunk_size={CHUNK_SIZE}")

    try:
        wb, ws, header_row, idx_map, headers = _open_workbook_rw(str(input_path))
    except Exception as e:
        msg = f"Erro ao abrir/interpretar arquivo Excel: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR open workbook error={str(e)[:1000]}")
        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": 0,
                "processed_success": 0,
                "failed_rows": 0,
                "execution_log_path": str(execution_log_path),
            },
        }

    total_rows = max(ws.max_row - header_row, 0)
    _append_execution_log(execution_log_path, f"INFO header_row={header_row}")
    _append_execution_log(execution_log_path, f"INFO total_rows_initial={total_rows}")

    try:
        while ws.max_row > header_row:
            chunk_number += 1
            start_row = header_row + 1
            end_row = min(start_row + CHUNK_SIZE - 1, ws.max_row)

            chunk_metrics = _process_chunk(
                ws=ws,
                idx_map=idx_map,
                headers=headers,
                start_row=start_row,
                end_row=end_row,
                file_path=str(input_path),
                failed_output_path=failed_output_path,
                processed_projects_by_ov=processed_projects_by_ov,
                execution_log_path=execution_log_path,
                chunk_number=chunk_number,
            )

            summary_metrics["processed_success"] += chunk_metrics["success"]
            summary_metrics["failed_rows"] += chunk_metrics["error"]
            summary_metrics["ignored_rows"] += chunk_metrics["ignored"]
            summary_metrics["projects_created"] += chunk_metrics["created"]
            summary_metrics["projects_updated"] += chunk_metrics["updated"]
            summary_metrics["warning_rows"] += chunk_metrics["warning_rows"]

            wb.save(str(input_path))

            remaining_rows_current = max(ws.max_row - header_row, 0)

            _append_execution_log(
                execution_log_path,
                f"INFO checkpoint_after_chunk chunk={chunk_number} success_total={summary_metrics['processed_success']} error_total={summary_metrics['failed_rows']} warning_rows_total={summary_metrics['warning_rows']} remaining_rows_current={remaining_rows_current}",
            )

            if chunk_number % IMPORT_CONTROL_PROGRESS_INTERVAL_CHUNKS == 0:
                _report_progress(
                    import_control_id=import_control_id,
                    total_rows=total_rows,
                    processed_success=summary_metrics["processed_success"],
                    failed_rows=summary_metrics["failed_rows"],
                    ignored_rows=summary_metrics["ignored_rows"],
                    remaining_rows=remaining_rows_current,
                    execution_log_path=execution_log_path,
                )

    except Exception as e:
        msg = f"Erro durante o processamento da importação: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR processing failed error={str(e)[:1000]}")

        try:
            wb.save(str(input_path))
        except Exception as save_ex:
            _append_execution_log(
                execution_log_path,
                f"ERROR failed saving workbook after exception error={str(save_ex)[:1000]}",
            )

        wb.close()

        finished_at = datetime.now()
        duration_seconds = int((finished_at - started_at).total_seconds())

        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": total_rows,
                "processed_success": summary_metrics["processed_success"],
                "failed_rows": summary_metrics["failed_rows"],
                "ignored_rows": summary_metrics["ignored_rows"],
                "projects_created": summary_metrics["projects_created"],
                "projects_updated": summary_metrics["projects_updated"],
                "warning_rows": summary_metrics["warning_rows"],
                "failed_file_path": str(failed_output_path),
                "execution_log_path": str(execution_log_path),
                "duration_seconds": duration_seconds,
            },
        }

    try:
        wb.save(str(input_path))
    except Exception as e:
        msg = f"Erro ao salvar arquivo Excel após processamento: {e}"
        _safe_log(str(input_path), 0, msg, None, None)
        _append_execution_log(execution_log_path, f"ERROR final workbook save failed error={str(e)[:1000]}")
        wb.close()

        finished_at = datetime.now()
        duration_seconds = int((finished_at - started_at).total_seconds())

        return {
            "status": "FAILED",
            "message": msg,
            "summary": {
                "file_path": str(input_path),
                "total_rows": total_rows,
                "processed_success": summary_metrics["processed_success"],
                "failed_rows": summary_metrics["failed_rows"],
                "ignored_rows": summary_metrics["ignored_rows"],
                "projects_created": summary_metrics["projects_created"],
                "projects_updated": summary_metrics["projects_updated"],
                "warning_rows": summary_metrics["warning_rows"],
                "failed_file_path": str(failed_output_path),
                "execution_log_path": str(execution_log_path),
                "duration_seconds": duration_seconds,
            },
        }

    wb.close()

    remaining_rows_in_input = 0
    try:
        wb_check = load_workbook(str(input_path), read_only=True, data_only=True)
        ws_check = wb_check.active
        remaining_rows_in_input = max(ws_check.max_row - header_row, 0)
        wb_check.close()
    except Exception as e:
        remaining_rows_in_input = -1
        _append_execution_log(
            execution_log_path,
            f"WARN failed checking remaining rows error={str(e)[:1000]}",
        )

    finished_at = datetime.now()
    duration_seconds = int((finished_at - started_at).total_seconds())

    status = "FAILED" if summary_metrics["failed_rows"] > 0 or summary_metrics["warning_rows"] > 0 else "FINISHED"

    msg = (
        f"{IMPORT_SOURCE} import concluído. "
        f"arquivo={str(input_path)}, total={total_rows}, "
        f"sucesso={summary_metrics['processed_success']}, "
        f"erros={summary_metrics['failed_rows']}, "
        f"advertencias={summary_metrics['warning_rows']}, "
        f"ignoradas={summary_metrics['ignored_rows']}, "
        f"projetos_criados={summary_metrics['projects_created']}, "
        f"projetos_atualizados={summary_metrics['projects_updated']}, "
        f"linhas_restantes_input={remaining_rows_in_input}, "
        f"arquivo_falhas={str(failed_output_path)}, "
        f"log_execucao={str(execution_log_path)}."
    )

    _append_execution_log(execution_log_path, f"INFO total_chunks={chunk_number}")
    _append_execution_log(execution_log_path, f"INFO total_rows_initial={total_rows}")
    _append_execution_log(execution_log_path, f"INFO processed_success={summary_metrics['processed_success']}")
    _append_execution_log(execution_log_path, f"INFO failed_rows={summary_metrics['failed_rows']}")
    _append_execution_log(execution_log_path, f"INFO warning_rows={summary_metrics['warning_rows']}")
    _append_execution_log(execution_log_path, f"INFO ignored_rows={summary_metrics['ignored_rows']}")
    _append_execution_log(execution_log_path, f"INFO projects_created={summary_metrics['projects_created']}")
    _append_execution_log(execution_log_path, f"INFO projects_updated={summary_metrics['projects_updated']}")
    _append_execution_log(execution_log_path, f"INFO remaining_rows_in_input={remaining_rows_in_input}")
    _append_execution_log(execution_log_path, f"INFO duration_seconds={duration_seconds}")
    _append_execution_log(
        execution_log_path,
        f"FINISH status={status} success={summary_metrics['processed_success']} error={summary_metrics['failed_rows']} warnings={summary_metrics['warning_rows']} ignored={summary_metrics['ignored_rows']} duration_seconds={duration_seconds}",
    )

    return {
        "status": status,
        "message": msg,
        "summary": {
            "file_path": str(input_path),
            "total_rows": total_rows,
            "processed_success": summary_metrics["processed_success"],
            "failed_rows": summary_metrics["failed_rows"],
            "warning_rows": summary_metrics["warning_rows"],
            "ignored_rows": summary_metrics["ignored_rows"],
            "projects_created": summary_metrics["projects_created"],
            "projects_updated": summary_metrics["projects_updated"],
            "failed_file_path": str(failed_output_path),
            "remaining_rows_in_input": remaining_rows_in_input,
            "execution_log_path": str(execution_log_path),
            "duration_seconds": duration_seconds,
            "total_chunks": chunk_number,
        },
    }
